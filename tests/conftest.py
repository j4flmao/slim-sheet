"""Fixtures are generated on the fly into tmp_path; no binary spreadsheets are committed."""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl
import pytest

from slim_sheet.config import Settings

ORDER_ROWS = 300


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return Settings(roots=(tmp_path.resolve(),), max_tokens=1500, max_cell_chars=40)


@pytest.fixture
def sales_xlsx(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    orders = wb.active
    assert orders is not None
    orders.title = "Orders"
    orders.append(["OrderID", "Date", "Region", "Qty", "Total"])
    for i in range(ORDER_ROWS):
        day = dt.datetime(2026, 1, 1) + dt.timedelta(days=i % 90)
        orders.append([10_000 + i, day, ["North", "South"][i % 2], i % 7 + 1, (i % 7 + 1) * 12.5])
    notes = wb.create_sheet("Notes")
    notes["B3"] = "x" * 500
    notes.merge_cells("D1:E1")
    notes["D1"] = "merged"
    hidden = wb.create_sheet("Secret")
    hidden.sheet_state = "hidden"
    hidden["A1"] = 1
    wb.create_sheet("Empty")
    path = tmp_path / "sales.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def people_csv(tmp_path: Path) -> Path:
    path = tmp_path / "people.csv"
    lines = ["name,age,city"] + [f"user{i},{20 + i % 40},City {i % 5}" for i in range(50)]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def _put(ws: object, top: int, left: int, rows: list[list[object]]) -> None:
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            ws.cell(row=top + r, column=left + c, value=value)  # type: ignore[attr-defined]


@pytest.fixture
def messy_xlsx(tmp_path: Path) -> Path:
    """Title row, then two side-by-side tables, plus a sheet with formulas."""
    wb = openpyxl.Workbook()
    report = wb.active
    assert report is not None
    report.title = "Report"
    report["A1"] = "Quarterly report"
    _put(
        report,
        3,
        1,
        [["Item", "Qty", "Price"], ["pen", 3, 1.5], ["ink", 10, 7.0], ["pad", 1, 2.25]],
    )
    _put(report, 3, 5, [["Code", "Note"], ["A", "refund issued"], ["B", "ok"]])
    calc = wb.create_sheet("Calc")
    calc["A1"], calc["A2"], calc["B1"] = 1, 2, "=SUM(A1:A2)"
    calc["C1"] = "=IFERROR(VLOOKUP(A1,Report!A3:C6,2,FALSE),0)"
    path = tmp_path / "messy.xlsx"
    wb.save(path)
    return path
