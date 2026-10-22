"""Design system: brands, inference, style_sheet, build_report, branded charts (PLAN §16)."""

from __future__ import annotations

import datetime as dt
import json
import zipfile
from dataclasses import replace
from pathlib import Path

import openpyxl
import pytest

from slim_sheet.adapters.brand import resolve_brand
from slim_sheet.adapters.writers.theme import has_custom_theme, theme_xml
from slim_sheet.config import Settings, load_settings
from slim_sheet.core.design import (
    BRANDS,
    DASH,
    brand_from_mapping,
    builtin_brand,
    formats,
    role_style,
    scaled_format,
)
from slim_sheet.core.errors import ErrorCode, SlimError
from slim_sheet.core.report_spec import (
    aggregate,
    choose_chart,
    infer_format,
    insight_title,
    looks_like_time,
    order,
    parse_chart,
    parse_kpi,
)
from slim_sheet.interfaces.handlers import Handlers
from slim_sheet.services import charts, design, report, structure

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]


@pytest.fixture
def shop_xlsx(tmp_path: Path) -> Path:
    """Sales!A1:F25: Month, Region, Product, Units, Revenue, Margin (fraction)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Sales"
    ws.append(["Month", "Region", "Product", "Units", "Revenue", "Margin"])
    regions = ["North", "South", "East", "West"]
    for m, month in enumerate(MONTHS):
        for r, region in enumerate(regions):
            ws.append(
                [month, region, f"P{r}", 10 + m + r, 1000.0 * (m + 1) + 250 * r, 0.1 + r / 100]
            )
    path = tmp_path / "shop.xlsx"
    wb.save(path)
    return path


def _xml(path: Path, part: str) -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read(part).decode("utf-8")


def _charts_xml(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        names = sorted(n for n in archive.namelist() if n.startswith("xl/charts/chart"))
        return [archive.read(n).decode("utf-8") for n in names]


# --- core: brands and formats ---------------------------------------------------------------


def test_finance_formats_use_parentheses_and_dash_for_zero() -> None:
    # Arrange
    brand = builtin_brand("slim")

    # Act
    named = formats(brand)

    # Assert
    assert named["integer"] == f'#,##0_);(#,##0);"{DASH}"_)'
    assert named["percent"].startswith("0.0%")
    assert named["growth"] == "+0.0%;-0.0%;0.0%"


def test_gov_brand_uses_minus_signs_and_long_dates() -> None:
    # Arrange
    brand = builtin_brand("gov")

    # Act
    named = formats(brand)

    # Assert
    assert named["integer"] == "#,##0"
    assert named["date"] == "d mmm yyyy"
    assert role_style("header", brand).fill is None  # government tables use no fills


def test_brand_file_layers_over_base_and_derives_palette() -> None:
    # Act
    brand = brand_from_mapping({"name": "Acme", "primary": "#c8102e", "currency": "$"}, "acme.json")

    # Assert
    assert brand.primary == "C8102E" and brand.palette[0] == "C8102E"
    assert brand.body_font == BRANDS["slim"].body_font
    assert formats(brand)["currency"].startswith('"$"#,##0')


@pytest.mark.parametrize(
    "data",
    [
        {},
        {"name": "X", "primary": "red"},
        {"name": "X", "colour": "#000000"},
        {"name": "X", "profile": "bank"},
    ],
)
def test_brand_file_errors_teach(data: dict[str, object]) -> None:
    # Act / Assert
    with pytest.raises(SlimError) as info:
        brand_from_mapping(data, "bad.json")
    assert info.value.code == ErrorCode.BAD_ARGUMENT and info.value.hint


def test_unknown_brand_and_role_list_choices() -> None:
    # Act / Assert
    with pytest.raises(SlimError, match="Unknown brand") as brand_err:
        builtin_brand("neon")
    with pytest.raises(SlimError, match="Unknown role") as role_err:
        role_style("hero", builtin_brand("slim"))
    assert "office" in (brand_err.value.hint or "") and "total" in (role_err.value.hint or "")


def test_scaled_format_compacts_large_headline_numbers() -> None:
    # Act / Assert
    assert scaled_format("integer", 2_500_000_000) == "bn"
    assert scaled_format("currency", 3_400_000) == "m"
    assert scaled_format("integer", 120_000) == "k"
    assert scaled_format("integer", 950) == "integer"
    assert scaled_format("percent", 5e9) == "percent"


# --- core: inference and specs --------------------------------------------------------------


@pytest.mark.parametrize(
    ("header", "values", "expected"),
    [
        ("Margin", [0.12, 0.3], "percent"),
        ("Revenue growth", [0.05, -0.02], "growth"),
        ("Conversion rate", [12.5, 30.1], "number1"),
        ("Year", [2024, 2025], "year"),
        ("Revenue", [1200.0, 5000.0], "integer"),
        ("Price", [1.25, 3.5], "number2"),
        ("Order ID", [1001, 1002], None),
        ("City", ["Hanoi", "Hue"], None),
        ("Day", [dt.date(2026, 1, 1)], "date"),
        ("Empty", [None, ""], None),
    ],
)
def test_infer_format_from_header_and_values(
    header: str, values: list[object], expected: str | None
) -> None:
    # Act
    key, reason = infer_format(header, values)

    # Assert
    assert key == expected and reason


def test_choose_chart_by_intent_and_shape() -> None:
    # Act / Assert
    assert choose_chart("auto", MONTHS, [1.0] * 6).kind == "line"
    assert choose_chart("auto", ["A", "B", "C"], [3.0, 2.0, 1.0]).kind == "column"
    assert choose_chart("compare", ["A very long category name", "B"], [1.0, 2.0]).kind == "bar"
    assert choose_chart("share", ["A", "B", "C"], [1.0, 2.0, 3.0]).kind == "doughnut"
    assert choose_chart("share", list("ABCDEFG"), [1.0] * 7).kind == "bar"
    assert choose_chart("rank", ["A", "B"], [1.0, 2.0]).sort == "desc"
    with pytest.raises(SlimError, match="intent"):
        choose_chart("pretty", ["A"], [1.0])


def test_looks_like_time_recognises_periods() -> None:
    # Act / Assert
    assert looks_like_time(["2024-01", "2024-02"])
    assert looks_like_time(["Q1 2025", "Q2 2025"])
    assert looks_like_time([2023, 2024, 2025])
    assert not looks_like_time(["North", "South"])
    assert not looks_like_time([])


def test_parse_kpi_variants() -> None:
    # Act
    total = parse_kpi("Revenue=sum:Revenue|currency")
    rows = parse_kpi("Orders=count")
    ratio = parse_kpi("Margin=ratio:Profit/Revenue")
    custom = parse_kpi("Custom==SUM(Sales!E2:E9)")
    bare = parse_kpi("Units")

    # Assert
    assert (total.agg, total.column, total.fmt) == ("sum", "Revenue", "currency")
    assert rows.agg == "count" and rows.column == ""
    assert (ratio.column, ratio.column2, ratio.fmt) == ("Profit", "Revenue", "percent")
    assert custom.agg == "formula" and custom.formula == "=SUM(Sales!E2:E9)"
    assert (bare.agg, bare.column) == ("sum", "Units")


@pytest.mark.parametrize("item", ["=sum:Revenue", "X=median:Revenue", "X=ratio:Profit", "X=sum"])
def test_parse_kpi_errors_include_examples(item: str) -> None:
    # Act / Assert
    with pytest.raises(SlimError) as info:
        parse_kpi(item)
    assert "e.g." in (info.value.hint or "") or "Use one of" in (info.value.hint or "")


def test_parse_chart_variants_and_errors() -> None:
    # Act
    plain = parse_chart("trend:Month:Revenue")
    counted = parse_chart("rank:Region:count:Orders by region")
    titled = parse_chart("share:Product:Units:avg:Mix: by product")

    # Assert
    assert (plain.intent, plain.category, plain.value, plain.agg) == (
        "trend",
        "Month",
        "Revenue",
        "sum",
    )
    assert (counted.value, counted.agg, counted.title) == ("count", "count", "Orders by region")
    assert titled.agg == "avg" and titled.title == "Mix: by product"
    with pytest.raises(SlimError, match="intent:Category:Value"):
        parse_chart("trend:Month")
    with pytest.raises(SlimError, match="unknown intent"):
        parse_chart("fancy:Month:Revenue")


def test_aggregate_order_and_insight_titles() -> None:
    # Arrange
    pairs = [("North", 10), ("South", 30), ("North", 20), (None, 99), ("East", "n/a")]

    # Act
    summed = aggregate(pairs, "sum")
    counted = aggregate(pairs, "count")
    ranked = order(summed, "desc")
    months = order([("Mar", 3.0), ("Jan", 1.0), ("Feb", 2.0)], "asc")

    # Assert
    assert summed == [("North", 30.0), ("South", 30.0), ("East", 0.0)]
    assert dict(counted)["North"] == 2.0
    assert ranked[0][1] == 30.0
    assert [m[0] for m in months] == ["Jan", "Feb", "Mar"]
    assert insight_title("trend", "Revenue", months) == "Revenue rose 200% from Jan to Mar"
    assert (
        insight_title("compare", "Units", [("A", 3.0), ("B", 1.0)])
        == "A leads Units with 75% of the total"
    )
    assert insight_title("compare", "P&L", [("A", -3.0), ("B", 1.0)]) == "B has the highest P&L"


# --- adapters: theme and brand resolution ---------------------------------------------------


def test_theme_xml_carries_brand_colours_and_fonts() -> None:
    # Act
    xml = theme_xml(builtin_brand("modern"))

    # Assert
    assert 'name="slim: modern"' in xml
    assert '<a:srgbClr val="156082"/>' in xml and '<a:srgbClr val="E97132"/>' in xml
    assert 'typeface="Aptos Display"' in xml and 'typeface="Aptos Narrow"' in xml
    assert "4F81BD" not in xml


def test_resolve_brand_order_argument_env_default(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    brand_file = tmp_path / "acme.json"
    brand_file.write_text(json.dumps({"name": "Acme", "primary": "#C8102E"}), encoding="utf-8")
    env_settings = replace(settings, brand=str(brand_file))

    # Act
    default, default_explicit = resolve_brand(None, settings)
    from_env, _ = resolve_brand("", env_settings)
    from_arg, arg_explicit = resolve_brand("mono", env_settings)

    # Assert
    assert default.name == "slim" and not default_explicit
    assert from_env.name == "Acme" and from_env.primary == "C8102E"
    assert from_arg.name == "mono" and arg_explicit


def test_brand_file_argument_follows_the_path_rules(
    tmp_path: Path, settings: Settings, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Arrange: a brand file in a data folder outside the roots, and one in a "system" folder
    from slim_sheet.adapters import sandbox

    monkeypatch.setattr(sandbox, "_system_folders", lambda: [(tmp_path / "sys").resolve()])
    home = tmp_path / "home"
    home.mkdir()
    settings = replace(settings, roots=(home.resolve(),))  # both folders below are outside it
    elsewhere = tmp_path / "brands"
    elsewhere.mkdir()
    (elsewhere / "acme.json").write_text('{"name": "Acme", "primary": "#C8102E"}', encoding="utf-8")
    (tmp_path / "sys").mkdir()
    (tmp_path / "sys" / "x.json").write_text('{"name": "X"}', encoding="utf-8")

    # Act
    allowed, explicit = resolve_brand(str(elsewhere / "acme.json"), settings)

    # Assert
    assert allowed.name == "Acme" and explicit
    with pytest.raises(SlimError) as system:
        resolve_brand(str(tmp_path / "sys" / "x.json"), settings)
    with pytest.raises(SlimError) as missing:
        resolve_brand("missing.json", settings)
    assert system.value.code == ErrorCode.PATH_DENIED
    assert missing.value.code == ErrorCode.FILE_NOT_FOUND


def test_brand_file_with_bad_json_is_reported(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    (tmp_path / "broken.json").write_text("{not json", encoding="utf-8")

    # Act / Assert
    with pytest.raises(SlimError, match="not valid JSON"):
        resolve_brand("broken.json", settings)


def test_settings_read_brand_from_env() -> None:
    # Act / Assert
    assert load_settings({"SLIM_SHEET_BRAND": "office"}).brand == "office"
    assert load_settings({"SLIM_SHEET_BRAND": "${user_config.brand}"}).brand is None


# --- services: create_workbook, style_sheet ---------------------------------------------------


def test_create_workbook_writes_brand_theme(tmp_path: Path, settings: Settings) -> None:
    # Act
    result = structure.create_workbook("new.xlsx", settings, ["Data"], brand="office")

    # Assert
    theme = _xml(tmp_path / "new.xlsx", "xl/theme/theme1.xml")
    assert 'name="slim: office"' in theme and "4472C4" in theme
    assert "brand office" in result.summary
    wb = openpyxl.load_workbook(tmp_path / "new.xlsx")
    assert "office header" in wb.named_styles


def test_style_sheet_report_infers_formats_and_styles_rows(
    shop_xlsx: Path, settings: Settings
) -> None:
    # Act
    result = design.style_sheet(str(shop_xlsx), "Sales", settings, overwrite=True)

    # Assert
    ws = openpyxl.load_workbook(shop_xlsx)["Sales"]
    assert "E Revenue: integer" in result.summary and "F Margin: percent" in result.summary
    assert ws["F2"].number_format.startswith("0.0%")
    assert ws["A1"].font.bold and ws["A1"].fill.fgColor.rgb == "FFF2F2F2"
    assert ws["E1"].alignment.horizontal == "right"  # numeric header follows its data
    assert ws.freeze_panes == "A2" and ws.sheet_view.showGridLines is False
    assert ws["B5"].font.name == BRANDS["slim"].body_font


def test_style_sheet_detects_total_row_and_data_filter(shop_xlsx: Path, settings: Settings) -> None:
    # Arrange
    wb = openpyxl.load_workbook(shop_xlsx)
    wb["Sales"].append(["Total", None, None, "=SUM(D2:D25)", "=SUM(E2:E25)", None])
    wb.save(shop_xlsx)

    # Act
    result = design.style_sheet(str(shop_xlsx), "Sales", settings, preset="data", overwrite=True)

    # Assert
    ws = openpyxl.load_workbook(shop_xlsx)["Sales"]
    assert "total row 26" in result.summary
    assert ws["E26"].font.bold and ws["E26"].border.top.style == "thin"
    assert ws["E26"].border.bottom.style == "double"  # finance profile
    assert ws.auto_filter.ref == "A1:F25"


def test_style_sheet_striped_minimal_and_forced_format(shop_xlsx: Path, settings: Settings) -> None:
    # Act
    design.style_sheet(str(shop_xlsx), "Sales", settings, preset="striped", out_path="s.xlsx")
    design.style_sheet(
        str(shop_xlsx),
        "Sales",
        settings,
        preset="minimal",
        number_format="number1",
        out_path="m.xlsx",
    )

    # Assert
    striped = openpyxl.load_workbook(shop_xlsx.with_name("s.xlsx"))["Sales"]
    minimal = openpyxl.load_workbook(shop_xlsx.with_name("m.xlsx"))["Sales"]
    assert striped["A3"].fill.fgColor.rgb == "FFF2F2F2" and striped["A2"].fill.fill_type is None
    assert minimal["A1"].fill.fill_type is None and minimal.freeze_panes is None
    assert minimal["E2"].number_format.startswith("#,##0.0")


def test_style_sheet_rejects_bad_preset_format_and_empty_sheet(
    sales_xlsx: Path, settings: Settings
) -> None:
    # Act / Assert
    with pytest.raises(SlimError, match="preset"):
        design.style_sheet(str(sales_xlsx), "Orders", settings, preset="fancy")
    with pytest.raises(SlimError, match="number_format"):
        design.style_sheet(str(sales_xlsx), "Orders", settings, number_format="money")
    with pytest.raises(SlimError, match="empty"):
        design.style_sheet(str(sales_xlsx), "Empty", settings)


def test_style_sheet_keeps_a_company_theme_unless_brand_is_explicit(
    shop_xlsx: Path, settings: Settings
) -> None:
    # Arrange: a workbook that carries someone else's theme
    wb = openpyxl.load_workbook(shop_xlsx)
    wb.loaded_theme = theme_xml(builtin_brand("office")).replace("slim: office", "Contoso").encode()
    wb.save(shop_xlsx)

    # Act
    kept = design.style_sheet(str(shop_xlsx), "Sales", settings, out_path="kept.xlsx")
    forced = design.style_sheet(
        str(shop_xlsx), "Sales", settings, brand="mono", out_path="forced.xlsx"
    )

    # Assert
    assert "kept theme" in kept.summary
    assert "Contoso" in _xml(shop_xlsx.with_name("kept.xlsx"), "xl/theme/theme1.xml")
    assert "slim: mono" in _xml(shop_xlsx.with_name("forced.xlsx"), "xl/theme/theme1.xml")
    assert "brand mono" in forced.summary
    assert has_custom_theme(openpyxl.load_workbook(shop_xlsx.with_name("kept.xlsx")))


# --- services: charts and build_report ------------------------------------------------------


def test_add_chart_auto_is_clean_and_keeps_axes_visible(
    shop_xlsx: Path, settings: Settings
) -> None:
    # Arrange
    wb = openpyxl.load_workbook(shop_xlsx)
    block = wb.create_sheet("Block")
    for row in [["Region", "Revenue"], ["North", 5.0], ["South", 9.0], ["East", 2.0]]:
        block.append(row)
    wb.save(shop_xlsx)

    # Act
    result = charts.add_chart(
        str(shop_xlsx), "Block", "auto", "A1:B4", settings, overwrite=True, highlight="South"
    )

    # Assert
    xml = _charts_xml(shop_xlsx)[0]
    assert "column chart" in result.summary and "auto:" in result.summary
    assert '<delete val="0"' in xml and '<roundedCorners val="0"' in xml
    assert "majorGridlines" not in xml and "<legend>" not in xml
    assert 'val="BFBFBF"' in xml and 'val="156082"' in xml  # grey context, one highlighted bar
    with pytest.raises(SlimError, match="highlight"):
        charts.add_chart(str(shop_xlsx), "Block", "column", "A1:B4", settings, highlight="Mars")


def test_build_report_creates_dashboard_tiles_charts_and_footer(
    shop_xlsx: Path, settings: Settings
) -> None:
    # Act
    result = report.build_report(
        str(shop_xlsx),
        "Sales",
        "Sales overview H1",
        settings,
        kpis=[
            "Revenue=sum:Revenue",
            "Orders=count",
            "Margin=ratio:Margin/Units",
            "Avg units=avg:Units",
        ],
        charts=[
            "trend:Month:Revenue",
            "rank:Region:Revenue",
            "share:Product:Units",
            "auto:Region:count",
        ],
        overwrite=True,
    )

    # Assert
    wb = openpyxl.load_workbook(shop_xlsx)
    board = wb["Dashboard"]
    assert wb.sheetnames[:2] == ["Dashboard", "Dashboard data"]
    assert board["B1"].value == "Sales overview H1" and board["B1"].font.size == 16
    assert (
        board["B5"].value == "=SUM('Sales'!$E$2:$E$25)"
        or board["B5"].value == "=SUM(Sales!$E$2:$E$25)"
    )
    assert board["D5"].value.startswith("=COUNTA(")
    assert board["F5"].value.startswith("=IFERROR(SUM(")
    assert board["B5"].number_format.startswith("#,##0_)")  # 93,000: below the k threshold
    assert board.sheet_view.showGridLines is False
    assert len(board._charts) == 4
    assert "rose" in result.summary and "leads Revenue" in result.summary
    assert "doughnut" in result.summary
    assert any("Source: Sales!A1:F25" in str(c.value) for row in board.iter_rows() for c in row)
    assert "slim: slim" in _xml(shop_xlsx, "xl/theme/theme1.xml")


def test_build_report_values_match_the_data(shop_xlsx: Path, settings: Settings) -> None:
    # Act
    result = report.build_report(
        str(shop_xlsx), "Sales", "T", settings, kpis=["Revenue=sum:Revenue"], overwrite=True
    )

    # Assert: 6 months x 4 regions: sum of 1000*(m+1) + 250*r
    expected = sum(1000.0 * (m + 1) + 250 * r for m in range(6) for r in range(4))
    assert f"Revenue = {expected:,.0f}" in result.summary


def test_build_report_rebuild_replaces_sheets_instead_of_duplicating(
    shop_xlsx: Path, settings: Settings
) -> None:
    # Act
    for _ in range(2):
        report.build_report(
            str(shop_xlsx), "Sales", "T", settings, charts=["rank:Region:Revenue"], overwrite=True
        )

    # Assert
    wb = openpyxl.load_workbook(shop_xlsx)
    assert wb.sheetnames == ["Dashboard", "Dashboard data", "Sales"]
    assert len(wb["Dashboard"]._charts) == 1


@pytest.mark.parametrize(
    ("kwargs", "match"),
    [
        ({}, "at least one KPI or chart"),
        ({"kpis": ["Sales=sum:Nope"]}, "No column 'Nope'"),
        ({"kpis": [f"K{i}=count" for i in range(7)]}, "at most"),
        ({"kpis": ["X=sum:Revenue|money"]}, "Unknown KPI format"),
        ({"charts": ["trend:Month:Revenue"], "sheet": "Sales"}, "replace the data sheet"),
    ],
)
def test_build_report_errors_teach(
    shop_xlsx: Path, settings: Settings, kwargs: dict[str, object], match: str
) -> None:
    # Act / Assert
    with pytest.raises(SlimError, match=match) as info:
        report.build_report(str(shop_xlsx), "Sales", "T", settings, **kwargs)  # type: ignore[arg-type]
    assert info.value.hint


def test_build_report_handler_and_format_range_role(shop_xlsx: Path, settings: Settings) -> None:
    # Arrange
    handlers = Handlers(settings)

    # Act
    built = handlers.build_report(
        str(shop_xlsx), "Sales", "Overview", kpis=["Units"], charts=["relationship:Units:Revenue"]
    )
    styled = handlers.format_range("shop.slim.xlsx", "Sales", "A1:F1", role="header")

    # Assert
    assert "built Dashboard" in built and "scatter" in built
    assert "role header" in styled
    ws = openpyxl.load_workbook(shop_xlsx.with_name("shop.slim.xlsx"))["Sales"]
    assert ws["A1"].font.bold


def test_autofit_measures_numbers_as_formatted(shop_xlsx: Path, settings: Settings) -> None:
    # Arrange: raw floats with many digits, shown as one-decimal percentages
    wb = openpyxl.load_workbook(shop_xlsx)
    ws = wb["Sales"]
    for row in range(2, 26):
        ws.cell(row=row, column=6).value = 0.123456789123
    wb.save(shop_xlsx)

    # Act
    design.style_sheet(str(shop_xlsx), "Sales", settings, overwrite=True)

    # Assert: sized for "12.3%" and the header, not for "0.123456789123"
    width = openpyxl.load_workbook(shop_xlsx)["Sales"].column_dimensions["F"].width
    assert width is not None and width <= 10
