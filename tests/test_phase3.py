from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from slim_sheet.config import Settings
from slim_sheet.core.diff import diff_formulas
from slim_sheet.core.errors import ErrorCode, SlimError
from slim_sheet.core.functions import FUNCTIONS, newer_than
from slim_sheet.core.refs import classify, max_extent, split_sheet, unquoted_sheet_with_space
from slim_sheet.core.validate import validate
from slim_sheet.interfaces.handlers import Handlers
from slim_sheet.services.describe import describe_cell
from slim_sheet.services.diff import diff
from slim_sheet.services.formula import validate_formula


@pytest.fixture
def model_xlsx(tmp_path: Path) -> Path:
    """A small financial model: formulas without cached values, a defined name, styles."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Calc"
    ws["A1"], ws["A2"], ws["A3"] = 10, 20, 30
    ws["B1"] = "=SUM(A1:A3)"
    ws["B2"] = "=B1*TaxRate"
    ws["B3"] = "=XLOOKUP(20,A1:A3,A1:A3)"
    ws["C1"], ws["C2"] = "=C2+1", "=C1+1"
    ws["D1"] = "=EDATE(DATE(2026,1,31),1)"
    ws["E1"] = "=FORMULATEXT(B1)"
    ws["B1"].font = Font(bold=True, color="FFFF0000")
    ws["B1"].fill = PatternFill("solid", fgColor="FFFFFF00")
    ws["B1"].number_format = "#,##0.00"
    ws["B1"].comment = Comment("Total of inputs", "analyst")
    dv = DataValidation(type="list", formula1='"yes,no"')
    ws.add_data_validation(dv)
    dv.add("F1")
    ws.merge_cells("G1:H2")
    rates = wb.create_sheet("Rate Card")
    rates["A1"] = 0.1
    wb.defined_names["TaxRate"] = DefinedName("TaxRate", attr_text="'Rate Card'!$A$1")
    path = tmp_path / "model.xlsx"
    wb.save(path)
    return path


# --- catalog -------------------------------------------------------------------------


def test_catalog_is_complete_and_consistent() -> None:
    # Act / Assert
    assert len(FUNCTIONS) >= 500
    for info in FUNCTIONS.values():
        assert info.min_args >= 0
        assert info.max_args == -1 or info.max_args >= info.min_args, info.name
        assert info.version in ("2003", "2007", "2010", "2013", "2016", "2019", "2021", "365")


@pytest.mark.parametrize(
    ("name", "target", "expected"),
    [("XLOOKUP", "2019", True), ("XLOOKUP", "365", False), ("IFERROR", "xls", True),
     ("SUM", "xls", False), ("CONCAT", "2016", True)],
)  # fmt: skip
def test_newer_than(name: str, target: str, expected: bool) -> None:
    # Act / Assert
    assert newer_than(FUNCTIONS[name], target) is expected


# --- refs ----------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "sheet", "kind"),
    [
        ("A1:B2", None, "cells"),
        ("'Rate Card'!$A$1", "Rate Card", "cells"),
        ("'It''s'!A1", "It's", "cells"),
        ("Data!A:C", "Data", "columns"),
        ("3:5", None, "rows"),
        ("Table1[Amount]", None, "table"),
        ("TaxRate", None, "name"),
        ("[Book2.xlsx]Sheet1!A1", "[Book2.xlsx]Sheet1", "external"),
    ],
)
def test_classify_refs(raw: str, sheet: str | None, kind: str) -> None:
    # Act
    ref = classify(raw)

    # Assert
    assert ref.sheet == sheet
    assert ref.kind == kind


def test_ref_helpers() -> None:
    # Act / Assert
    assert split_sheet("A1") == (None, "A1")
    assert max_extent(classify("A1:C70000")) == (70000, 2)
    assert unquoted_sheet_with_space("=SUM(My Data!A1:A3)") == "My Data"
    assert unquoted_sheet_with_space("=SUM('My Data'!A1:A3)") is None
    assert unquoted_sheet_with_space('="a b!c"') is None


# --- validate (pure) -----------------------------------------------------------------


@pytest.mark.parametrize(
    ("formula", "target", "fragment", "fix"),
    [
        ("=XLOOKUP(A1,B:B,C:C)", "2016", "needs Excel 2021", "INDEX("),
        ("=IFERROR(A1,0)", "xls", "needs Excel 2007", "IF(ISERROR"),
        ("=SUM(A1:A70000)", "xls", "outside the xls grid", None),
        ("=SUMM(A1,A2)", "365", "not an Excel function", "rename to SUM"),
        ("=SUM(A1;A2)", "365", "';'", "use ','"),
        ('=WEBSERVICE("http://x")', "365", "reaches the network", None),
        ("=[Other.xlsx]S!A1", "365", "another workbook", None),
    ],
)
def test_validate_reports_errors_and_fixes(
    formula: str, target: str, fragment: str, fix: str | None
) -> None:
    # Act
    report = validate(formula, target)

    # Assert
    assert not report.ok
    assert any(fragment in e for e in report.errors), report.errors
    if fix:
        assert any(fix in f for f in report.fixes), report.fixes


def test_validate_against_workbook_names_and_sheets() -> None:
    # Act
    good = validate("=TaxRate*'Rate Card'!A1", sheets=["Calc", "Rate Card"], names=["TaxRate"])
    bad = validate("=TaxRat*Rates!A1", sheets=["Calc", "Rate Card"], names=["TaxRate"])

    # Assert
    assert good.ok
    assert len(bad.errors) == 2
    assert any("TaxRate" in f for f in bad.fixes)


def test_validate_rejects_unknown_target() -> None:
    # Act / Assert
    with pytest.raises(SlimError) as info:
        validate("=1", "1999")
    assert "365" in info.value.hint


# --- validate (service) --------------------------------------------------------------


def test_validate_formula_reads_cell_and_workbook(model_xlsx: Path, settings: Settings) -> None:
    # Act
    report, address = validate_formula(settings, path="model.xlsx", cell="Calc!B3", target="2019")
    named, _ = validate_formula(settings, formula="=B1*TaxRate", path="model.xlsx")

    # Assert
    assert address == "Calc!B3"
    assert any("XLOOKUP" in e for e in report.errors)
    assert named.ok


# --- evaluate ------------------------------------------------------------------------

formulas_lib = pytest.importorskip("formulas")


def test_evaluate_chases_uncached_precedents(model_xlsx: Path, settings: Settings) -> None:
    # Arrange
    from slim_sheet.services.evaluate import evaluate

    # Act
    result = evaluate(settings, path="model.xlsx", cell="Calc!B2")

    # Assert
    assert result.value == pytest.approx(6.0)
    assert result.recomputed == ("Calc!B1",)
    assert ("Rate Card!A1", "= 0.1") in result.inputs


@pytest.mark.parametrize(
    ("kwargs", "expected"),
    [
        ({"formula": "=ROUND(2.345,2)+1"}, 3.35),
        ({"path": "model.xlsx", "formula": "=SUM(A:A)", "sheet": "Calc"}, 60),
        ({"path": "model.xlsx", "formula": "='Rate Card'!A1*100"}, 10),
        ({"path": "model.xlsx", "cell": "B3", "sheet": "Calc"}, 20),
    ],
)
def test_evaluate_values(
    model_xlsx: Path, settings: Settings, kwargs: dict[str, str], expected: float
) -> None:
    # Arrange
    from slim_sheet.services.evaluate import evaluate

    # Act
    result = evaluate(settings, **kwargs)

    # Assert
    assert result.value == pytest.approx(expected)


def test_evaluate_date_note_and_value_cell(model_xlsx: Path, settings: Settings) -> None:
    # Arrange
    from slim_sheet.services.evaluate import evaluate

    # Act
    dated = evaluate(settings, path="model.xlsx", cell="Calc!D1")
    plain = evaluate(settings, path="model.xlsx", cell="Calc!A1")

    # Assert
    assert "as a date: 2026-02-28" in dated.notes
    assert plain.value == 10
    assert "not a formula" in plain.notes[0]


@pytest.mark.parametrize(
    ("kwargs", "code"),
    [
        ({"path": "model.xlsx", "cell": "Calc!C1"}, ErrorCode.CIRCULAR_REFERENCE),
        ({"formula": "=A1*2"}, ErrorCode.BAD_ARGUMENT),
        ({"formula": "=SUM((1)"}, ErrorCode.FORMULA_SYNTAX),
        ({"path": "model.xlsx", "formula": "=Table1[Col]"}, ErrorCode.UNSUPPORTED_FUNCTION),
        ({}, ErrorCode.BAD_ARGUMENT),
    ],
)
def test_evaluate_errors(
    model_xlsx: Path, settings: Settings, kwargs: dict[str, str], code: ErrorCode
) -> None:
    # Arrange
    from slim_sheet.services.evaluate import evaluate

    # Act / Assert
    with pytest.raises(SlimError) as info:
        evaluate(settings, **kwargs)
    assert info.value.code == code


def test_evaluate_flags_stale_cache(tmp_path: Path, settings: Settings) -> None:
    # Arrange: a cached value that no longer matches what the inputs produce
    from slim_sheet.services.evaluate import _Evaluator, _result

    # Act
    result = _result(_Evaluator(None, None), "=1+1", "S!A1", 2, 3)

    # Assert
    assert any("differs from Excel's cached value 3" in n for n in result.notes)


# --- describe_cell -------------------------------------------------------------------


def test_describe_cell_reports_style_and_context(model_xlsx: Path, settings: Settings) -> None:
    # Act
    facts = dict(describe_cell(settings, "model.xlsx", "Calc!B1").facts)
    validation = dict(describe_cell(settings, "model.xlsx", "F1", sheet="Calc").facts)
    merged = dict(describe_cell(settings, "model.xlsx", "Calc!H2").facts)

    # Assert
    assert facts["formula"] == "=SUM(A1:A3)"
    assert facts["number format"] == "#,##0.00"
    assert "bold" in facts["font"] and "#FF0000" in facts["font"]
    assert facts["fill"] == "#FFFF00"
    assert facts["comment"].startswith("Total of inputs")
    assert validation["validation"] == 'list: "yes,no"'
    assert merged["merged"] == "G1:H2"


def test_describe_cell_rejects_ranges_and_handles_csv(
    model_xlsx: Path, people_csv: Path, settings: Settings
) -> None:
    # Act
    csv_cell = describe_cell(settings, "people.csv", "A2")

    # Assert
    assert csv_cell.value == "user0"
    assert "values only" in csv_cell.facts[0][1]
    with pytest.raises(SlimError) as info:
        describe_cell(settings, "model.xlsx", "A1:B2")
    assert info.value.code == ErrorCode.BAD_ARGUMENT


# --- formula diff --------------------------------------------------------------------


def test_diff_formulas_ignores_case_and_spaces() -> None:
    # Act
    changes = diff_formulas(
        {(0, 0): "=sum(A1:A3)", (1, 0): '=A1&" x "', (2, 0): "=A1"},
        {(0, 0): "=SUM( A1:A3 )", (1, 0): '=A1&"  x "', (3, 0): "=B1"},
    )

    # Assert
    assert [(r, c) for r, c, *_ in changes] == [(1, 0), (2, 0), (3, 0)]


def test_diff_reports_formula_changes_between_xlsx(model_xlsx: Path, settings: Settings) -> None:
    # Arrange
    copy = model_xlsx.with_name("model2.xlsx")
    wb = openpyxl.load_workbook(model_xlsx)
    wb["Calc"]["B1"] = "=SUM(A1:A2)"
    wb.save(copy)

    # Act
    result = diff("model.xlsx", "model2.xlsx", settings)

    # Assert
    assert ("Calc!B1 (formula)", "=SUM(A1:A3)", "=SUM(A1:A2)") in result.changes
    assert "Calc: 1 changed formula(s)" in result.summary


# --- interfaces ----------------------------------------------------------------------


def test_phase3_handlers_render(model_xlsx: Path, settings: Settings) -> None:
    # Arrange
    h = Handlers(settings)

    # Act
    validated = h.evaluate(formula="=XLOOKUP(1,A:A,B:B)", check_only=True, target="2016")
    evaluated = h.evaluate(path="model.xlsx", cell="Calc!B2")
    described = h.describe_cell("model.xlsx", "Calc!B1")

    # Assert
    assert "target Excel 2016: 1 problem(s)" in validated
    assert "fix: replace XLOOKUP" in validated
    assert "result: 6" in evaluated
    assert "recomputed (no cached value): Calc!B1" in evaluated
    assert "comment: Total of inputs (by analyst)" in described


def test_write_warns_about_broken_formulas(sales_xlsx: Path, settings: Settings) -> None:
    # Arrange
    from slim_sheet.services.write import WriteRequest, write

    # Act
    result = write(
        WriteRequest("sales.xlsx", "Orders", "H2", [["=SUMM(A2:A3)", "=SUM(A2:A3)"]]), settings
    )

    # Assert
    assert any(
        w.startswith("H2 =SUMM(A2:A3): SUMM is not an Excel function") for w in result.warnings
    )
    assert not any(w.startswith("I2") for w in result.warnings)


def test_evaluate_without_formula_extra_explains_install(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # Arrange
    import sys

    from slim_sheet.adapters import formula_engine

    formula_engine._lib.cache_clear()
    monkeypatch.setitem(sys.modules, "formulas", None)

    # Act / Assert
    try:
        with pytest.raises(SlimError) as info:
            formula_engine.compile_formula("=1+1")
        assert info.value.code == ErrorCode.MISSING_EXTRA
        assert "slim-sheet[formula]" in info.value.hint
    finally:
        formula_engine._lib.cache_clear()
