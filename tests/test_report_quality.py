"""Client-ready report rules: KPI variance notes, key messages, ranked tables, print setup and
chart details (titles, formats, negatives, thin slices). docs/prompts/design.md §4."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import openpyxl
import pytest

from slim_sheet.adapters.writers.chart_fixups import fix_chart_xml
from slim_sheet.adapters.writers.chart_style import axis_format, clean_title
from slim_sheet.config import Settings
from slim_sheet.core.errors import ErrorCode, SlimError
from slim_sheet.core.report_spec import infer_format, insight_title, parse_kpi
from slim_sheet.interfaces.handlers import Handlers
from slim_sheet.services import charts, design, report

REGIONS = ("North", "South", "East", "West", "Islands")


@pytest.fixture
def two_years(tmp_path: Path) -> Path:
    """Sales!A1:F21: Region, Product, Revenue 2025, Revenue 2024, Cost, Growth (one negative)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Sales"
    ws.append(["Region", "Product", "Revenue 2025", "Revenue 2024", "Cost", "Growth"])
    for r, region in enumerate(REGIONS):
        for p in range(4):
            now, before = 1000.0 + 400 * r + 50 * p, 900.0 + 300 * r + 40 * p
            ws.append([region, f"P{p}", now, before, now * 0.6, now / before - 1])
    ws["F2"] = -0.25
    path = tmp_path / "sales.xlsx"
    wb.save(path)
    return path


def _charts_xml(path: Path) -> list[str]:
    """Chart parts with the chart namespace prefix and self-closing spaces removed."""
    with zipfile.ZipFile(path) as archive:
        names = sorted(n for n in archive.namelist() if n.startswith("xl/charts/chart"))
        parts = [archive.read(n).decode("utf-8") for n in names]
    return [re.sub(r"<(/?)c:", r"<", x).replace(" />", "/>") for x in parts]


def _cells(ws: object) -> list[object]:
    return [c.value for row in ws.iter_rows() for c in row]  # type: ignore[attr-defined]


# --- KPI grammar ----------------------------------------------------------------------------


def test_parse_kpi_reads_format_and_comparison_in_any_order() -> None:
    # Act
    first = parse_kpi("Revenue=sum:Revenue 2025|vs:Revenue 2024:2024|currency")
    ratio = parse_kpi("Margin=ratio:A/B|vs:C/D")

    # Assert
    assert (first.column, first.fmt, first.vs, first.vs_label) == (
        "Revenue 2025",
        "currency",
        "Revenue 2024",
        "2024",
    )
    assert (ratio.vs, ratio.vs_label, ratio.fmt) == ("C/D", "prior", "percent")


@pytest.mark.parametrize(
    "item",
    [
        "Orders=count|vs:Orders 2024",  # a count has nothing to compare
        "Margin=ratio:A/B|vs:C",  # a ratio compares with a ratio
        "Revenue=sum:R|vs:A/B",  # a value compares with a column
        "Revenue=sum:R|vs:",  # no column
        "Revenue=sum:R|currency|percent",  # two formats
    ],
)
def test_parse_kpi_rejects_bad_comparisons(item: str) -> None:
    # Act / Assert
    with pytest.raises(SlimError) as caught:
        parse_kpi(item)
    assert caught.value.code == ErrorCode.BAD_ARGUMENT
    assert caught.value.hint


def test_insight_titles_never_call_an_average_rate_a_share_of_the_total() -> None:
    # Arrange
    points = [("Real estate", 0.45), ("Banks", 0.12)]

    # Act
    rate = insight_title("rank", "Profit growth", points, "avg", "growth")
    summed = insight_title("rank", "Profit", [("A", 300.0), ("B", 100.0)], "sum", "number1")

    # Assert
    assert rate == "Real estate has the highest Profit growth (+45.0%)"
    assert summed == "A leads Profit with 75% of the total"


# --- dashboard ------------------------------------------------------------------------------


def test_kpi_comparison_is_a_live_coloured_note(two_years: Path, settings: Settings) -> None:
    # Act
    result = report.build_report(
        str(two_years),
        "Sales",
        "Sales 2025",
        settings,
        kpis=[
            "Revenue=sum:Revenue 2025|vs:Revenue 2024:2024",
            "Cost=sum:Cost|vs:Revenue 2024",
            "Margin=ratio:Cost/Revenue 2025|vs:Cost/Revenue 2024:2024",
        ],
        overwrite=True,
    )

    # Assert
    board = openpyxl.load_workbook(two_years)["Dashboard"]
    assert board["B6"].value.startswith("=IFERROR(TEXT($B$5/(SUM(")
    assert board["B6"].value.endswith('&" vs 2024","n/a")')
    assert '" pts vs 2024"' in board["F6"].value
    rules = {str(r.sqref): [x.formula[0] for x in r.rules] for r in board.conditional_formatting}
    assert ">=(SUM(" in rules["B6"][0]  # revenue up is good
    assert "<=(SUM(" in rules["D6"][0]  # cost up is bad
    revenue_now = sum(1000.0 + 400 * r + 50 * p for r in range(5) for p in range(4))
    revenue_before = sum(900.0 + 300 * r + 40 * p for r in range(5) for p in range(4))
    assert f"{revenue_now / revenue_before - 1:+.1%} vs 2024" in result.summary


def test_key_messages_sit_between_tiles_and_charts(two_years: Path, settings: Settings) -> None:
    # Act
    result = report.build_report(
        str(two_years),
        "Sales",
        "Sales 2025",
        settings,
        kpis=["Revenue=sum:Revenue 2025"],
        charts=["rank:Region:Revenue 2025"],
        messages=["Islands grew fastest.", "  ", "Cost stayed at 60% of revenue."],
        overwrite=True,
    )

    # Assert
    board = openpyxl.load_workbook(two_years)["Dashboard"]
    assert board["B8"].value == "Key messages"
    assert board["B9"].value.endswith("Islands grew fastest.")
    assert board["B10"].value.endswith("Cost stayed at 60% of revenue.")
    assert "2 key message(s)" in result.summary
    assert "chart B12:" in result.summary  # the grid starts below the messages


def test_too_many_key_messages_are_refused(two_years: Path, settings: Settings) -> None:
    # Act / Assert
    with pytest.raises(SlimError) as caught:
        report.build_report(
            str(two_years), "Sales", "T", settings, kpis=["Revenue=sum:Revenue 2025"],
            messages=[f"Finding {i}" for i in range(6)], overwrite=True,
        )  # fmt: skip
    assert caught.value.code == ErrorCode.TOO_LARGE


def test_table_slot_is_a_live_ranked_table(two_years: Path, settings: Settings) -> None:
    # Act
    result = report.build_report(
        str(two_years),
        "Sales",
        "T",
        settings,
        charts=["table:Region:Revenue 2025", "table:Region:Growth:avg:Growth by region"],
        overwrite=True,
    )

    # Assert
    board = openpyxl.load_workbook(two_years)["Dashboard"]
    assert board["B8"].value == "Top 5 Region by Revenue 2025"
    assert board["B10"].value == "Islands"  # largest first
    assert board["D10"].value == "=SUMIFS(Sales!$C$2:$C$21,Sales!$A$2:$A$21,$B$10)"
    assert board["F9"].value == "Share" and board["F10"].number_format == "0.0%"
    assert board["B15"].value == "Total (all)"
    assert board["H8"].value == "Growth by region"
    assert board["J10"].value.startswith("=AVERAGEIFS(")
    assert board["L9"].value is None  # averages are not shares of a whole
    assert "table B8: top 5 of 5 Region" in result.summary
    assert len(board._charts) == 0


def test_dashboard_prints_on_one_landscape_page_without_tool_branding(
    two_years: Path, settings: Settings
) -> None:
    # Act
    report.build_report(
        str(two_years), "Sales", "Sales 2025", settings,
        charts=["rank:Region:Revenue 2025"], overwrite=True,
    )  # fmt: skip

    # Assert
    wb = openpyxl.load_workbook(two_years)
    board = wb["Dashboard"]
    assert board.page_setup.orientation == "landscape"
    assert (board.page_setup.fitToWidth, board.page_setup.fitToHeight) == (1, 1)
    assert board.print_area == "'Dashboard'!$A$1:$M$25"
    assert board.oddFooter.right.text == "Page &P of &N"
    assert board.oddFooter.left.text == "Sales 2025"
    assert not any("slim-sheet" in str(v) for v in _cells(board) if v)
    assert wb.sheetnames[-1] == "Dashboard data"
    assert wb["Dashboard data"].sheet_properties.tabColor.rgb == "FFBFBFBF"


# --- tables ---------------------------------------------------------------------------------


def test_style_sheet_makes_a_wide_table_print_ready(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Wide"
    ws.append([f"A long column header {i}" for i in range(12)])
    ws.append([1000 * i for i in range(12)])
    wb.save(tmp_path / "w.xlsx")

    # Act
    result = design.style_sheet(str(tmp_path / "w.xlsx"), "Wide", settings, overwrite=True)

    # Assert
    sheet = openpyxl.load_workbook(tmp_path / "w.xlsx")["Wide"]
    assert "prints landscape, fit to width, header repeated" in result.summary
    assert sheet.page_setup.orientation == "landscape"
    assert (sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight) == (1, 0)
    assert sheet.print_title_rows == "$1:$1"


def test_style_sheet_keeps_the_authors_print_area(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Item", "Amount"])
    ws.append(["a", 1])
    ws.print_area = "A1:A2"
    wb.save(tmp_path / "p.xlsx")

    # Act
    result = design.style_sheet(str(tmp_path / "p.xlsx"), "Sheet", settings, overwrite=True)

    # Assert
    sheet = openpyxl.load_workbook(tmp_path / "p.xlsx")["Sheet"]
    assert sheet.print_area == "'Sheet'!$A$1:$A$2"
    assert "prints" not in result.summary


# --- charts ---------------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "expected"),
    [(None, None), ("None", None), (" null ", None), ("", None), ("n/a", None), (" Tax ", "Tax")],
)
def test_clean_title_drops_placeholder_words(raw: str | None, expected: str | None) -> None:
    # Act / Assert
    assert clean_title(raw) == expected


def _block(tmp_path: Path, rows: list[list[object]], fmt: str | None = None) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "S"
    for row in rows:
        ws.append(row)
    if fmt:
        for (cell,) in ws.iter_rows(min_row=2, min_col=2, max_col=2):
            cell.number_format = fmt
    path = tmp_path / "c.xlsx"
    wb.save(path)
    return path


def test_add_chart_never_prints_placeholder_titles(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    path = _block(tmp_path, [["Region", "Profit"], ["A", 3], ["B", 2]])

    # Act
    charts.add_chart(str(path), "S", "column", "A1:B3", settings, title="None",
                     x_title="null", y_title="USD bn", overwrite=True)  # fmt: skip

    # Assert
    xml = _charts_xml(path)[0]
    assert ">None<" not in xml and ">null<" not in xml
    assert ">USD bn<" in xml


def test_bar_axis_titles_follow_what_the_reader_sees(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    path = _block(tmp_path, [["Region", "Profit"], ["A", 3], ["B", 2]])

    # Act: on horizontal bars the horizontal axis is the value axis
    charts.add_chart(str(path), "S", "bar", "A1:B3", settings, x_title="USD bn", overwrite=True)

    # Assert
    xml = _charts_xml(path)[0]
    value_axis = xml[xml.index("<valAx>") : xml.index("</valAx>")]
    assert ">USD bn<" in value_axis
    assert '<delete val="0"/>' in value_axis  # a titled value axis stays visible


def test_chart_uses_the_cells_number_format_and_clears_negatives(
    tmp_path: Path, settings: Settings
) -> None:
    # Arrange
    path = _block(tmp_path, [["Region", "Growth"], ["A", 0.12], ["B", -0.05]], "0.0%")

    # Act
    charts.add_chart(str(path), "S", "column", "A1:B3", settings, overwrite=True)

    # Assert
    xml = _charts_xml(path)[0]
    assert 'formatCode="0.0%"' in xml
    category_axis = xml[xml.index("<catAx>") : xml.index("</catAx>")]
    assert '<tickLblPos val="low"/>' in category_axis
    value_axis = xml[xml.index("<valAx>") : xml.index("</valAx>")]
    assert '<delete val="1"/>' in value_axis  # labelled bars need no value axis


def test_chart_infers_a_format_from_the_header(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    path = _block(tmp_path, [["Region", "Net margin"], ["A", 0.2], ["B", 0.31]])

    # Act
    charts.add_chart(str(path), "S", "column", "A1:B3", settings, overwrite=True)

    # Assert
    assert re.search(r'formatCode="0\.0%[^"]*"', _charts_xml(path)[0])


def test_thin_pie_slices_get_no_label(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    rows: list[list[object]] = [["Part", "Value"], ["Big", 97], ["Mid", 50], ["Tiny", 1]]
    path = _block(tmp_path, rows)

    # Act
    charts.add_chart(str(path), "S", "pie", "A1:B4", settings, overwrite=True)

    # Assert
    xml = _charts_xml(path)[0]
    tiny = re.search(r'<dLbl><idx val="2"/>.*?</dLbl>', xml)
    assert tiny and '<showPercent val="0"/>' in tiny.group(0)
    assert '<idx val="0"/><spPr>' not in xml.split("<dLbls>")[1].split("</dLbls>")[0]


def test_doughnut_uses_a_legend_and_contrasting_ring_labels(
    tmp_path: Path, settings: Settings
) -> None:
    # Arrange
    rows: list[list[object]] = [["Part", "Value"], ["A", 60], ["B", 40]]
    path = _block(tmp_path, rows)

    # Act
    charts.add_chart(str(path), "S", "doughnut", "A1:B3", settings, overwrite=True)

    # Assert
    xml = _charts_xml(path)[0]
    assert '<legendPos val="r"/>' in xml
    assert 'srgbClr val="FFFFFF"' in xml.split("<dLbls>")[1]  # white on the dark brand slice


def test_a_single_trend_line_labels_its_latest_value(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    rows: list[list[object]] = [["Month", "Revenue"]] + [[f"M{i}", 100 + i] for i in range(6)]
    path = _block(tmp_path, rows)

    # Act
    charts.add_chart(str(path), "S", "line", "A1:B7", settings, overwrite=True)

    # Assert
    xml = _charts_xml(path)[0]
    assert re.search(r'<dLbl><idx val="5"/>.*?<showVal val="1"/>', xml)


def test_add_chart_refuses_a_text_column_as_a_series(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    path = _block(tmp_path, [["Period", "Market", "Revenue"], ["FY1", "US", 3], ["FY2", "US", 4]])

    # Act / Assert
    with pytest.raises(SlimError) as caught:
        charts.add_chart(str(path), "S", "column", "A1:C3", settings, overwrite=True)
    assert caught.value.code == ErrorCode.BAD_ARGUMENT
    assert "Column B ('Market') holds text" in caught.value.message
    assert caught.value.hint


@pytest.mark.parametrize(
    ("fmt", "values", "expected"),
    [
        ("#,##0.0_);(#,##0.0)", (4500.0, 12.0), "#,##0_);(#,##0)"),
        ("0.0%", (0.25, 0.4), "0%"),
        ("0.0%", (0.02, 0.05), "0.0%"),  # a small range keeps its decimals
        ("#,##0.00", (3.2,), "#,##0.00"),
        ("#,##0.0", (), "#,##0.0"),  # no values known: unchanged
    ],
)
def test_axis_format_drops_decimals_ticks_never_need(
    fmt: str, values: tuple[float, ...], expected: str
) -> None:
    # Act / Assert
    assert axis_format(fmt, values) == expected


def test_a_very_wide_table_prints_across_pages_at_a_readable_size(
    tmp_path: Path, settings: Settings
) -> None:
    # Arrange
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Data"
    ws.append([f"A rather long column header number {i}" for i in range(30)])
    ws.append(list(range(30)))
    wb.save(tmp_path / "d.xlsx")

    # Act
    result = design.style_sheet(str(tmp_path / "d.xlsx"), "Data", settings, overwrite=True)

    # Assert
    sheet = openpyxl.load_workbook(tmp_path / "d.xlsx")["Data"]
    assert "too wide for one page" in result.summary
    assert sheet.page_setup.orientation == "landscape"
    assert (
        sheet.sheet_properties.pageSetUpPr is None
        or not sheet.sheet_properties.pageSetUpPr.fitToPage
    )
    assert sheet.print_title_cols == "$A:$A"


# --- existing charts, Vietnamese headers, openpyxl chart repairs -------------------------------


@pytest.mark.parametrize(
    ("header", "values", "expected"),
    [
        ("Tăng trưởng LN", [0.12, -0.05], "growth"),
        ("TT DT", [0.3, 0.1], "growth"),
        ("Biên lợi nhuận", [0.2, 0.31], "percent"),
        ("Tỷ trọng lũy kế", [0.5, 0.9], "percent"),
        ("LN tăng thêm", [12.5, 30.25], "number1"),  # an amount added, not a growth rate
        ("Năm thành lập", [1999, 2004], "year"),
        ("Mã", [101, 205], None),
    ],
)
def test_vietnamese_headers_pick_the_same_formats_as_english(
    header: str, values: list[float], expected: str | None
) -> None:
    # Act / Assert
    assert infer_format(header, values)[0] == expected


def test_fix_chart_xml_links_label_formats_and_drops_none_runs() -> None:
    # Arrange
    xml = (
        '<c:dLbls><c:numFmt formatCode="0.0%"/><c:txPr><a:bodyPr/><a:p><a:r><a:t>None</a:t>'
        '</a:r><a:endParaRPr/></a:p></c:txPr></c:dLbls><c:valAx><c:numFmt formatCode="0" '
        'sourceLinked="1"/></c:valAx><c:title><c:tx><c:rich><a:p><a:r><a:t>Keep</a:t></a:r>'
        "</a:p></c:rich></c:tx></c:title>"
    )

    # Act
    fixed = fix_chart_xml(xml)

    # Assert
    assert '<c:numFmt formatCode="0.0%" sourceLinked="0"/>' in fixed
    assert 'formatCode="0" sourceLinked="1"' in fixed  # an explicit link is left alone
    assert ">None<" not in fixed
    assert "<a:t>Keep</a:t>" in fixed  # title text is not txPr


def test_label_formats_survive_later_edits(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    path = _block(tmp_path, [["Region", "Growth"], ["A", 0.12], ["B", 0.3]])
    charts.add_chart(str(path), "S", "column", "A1:B3", settings, overwrite=True)

    # Act: two more load/save cycles through other edits
    for cell in ("D1", "D2"):
        Handlers(settings).write("c.xlsx", "S", cell, [["x"]], overwrite=True)

    # Assert
    xml = _charts_xml(path)[0]
    labels = xml[xml.index("<dLbls>") : xml.index("</dLbls>")]
    assert 'sourceLinked="0"' in labels
    assert ">None<" not in xml


def _old_style_chart(tmp_path: Path) -> Path:
    """A chart as another tool leaves it: default style, raw numbers, a 'None' axis title."""
    from openpyxl.chart import BarChart, Reference

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "S"
    for row in [["Market", "Tăng trưởng LN"], ["US", 0.288], ["India", -0.089], ["VN", 0.263]]:
        ws.append(row)
    chart = BarChart()
    chart.type = "bar"
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
    chart.x_axis.title = "None"
    chart.y_axis.scaling.max = 0.3
    ws.add_chart(chart, "D2")
    path = tmp_path / "old.xlsx"
    wb.save(path)
    return path


def test_style_sheet_restyles_existing_charts(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    path = _old_style_chart(tmp_path)

    # Act
    result = design.style_sheet(str(path), "S", settings, preset="charts", overwrite=True)

    # Assert
    xml = _charts_xml(path)[0]
    assert "S (brand slim): restyled 1 chart(s)" in result.summary
    assert 'formatCode="+0.0%;-0.0%;0.0%" sourceLinked="0"' in xml  # from the Vietnamese header
    assert ">None<" not in xml
    assert '<invertIfNegative val="0"/>' in xml  # negative bars stay solid
    assert "<max " not in xml  # the old fixed maximum would clip the end labels
    assert openpyxl.load_workbook(path)["S"]["A1"].font.b is False  # cells untouched


def test_style_sheet_reports_kinds_it_leaves_alone(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    from openpyxl.chart import RadarChart, Reference

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "S"
    for row in [["Axis", "Score"], ["a", 1], ["b", 2], ["c", 3]]:
        ws.append(row)
    radar = RadarChart()
    radar.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    ws.add_chart(radar, "D2")
    wb.save(tmp_path / "r.xlsx")

    # Act
    result = design.style_sheet(str(tmp_path / "r.xlsx"), "S", settings, overwrite=True)

    # Assert
    assert "left 1 chart(s) as they are (RadarChart)" in result.summary


@pytest.mark.parametrize(
    ("values", "expected"),
    [
        ([25054.0, 7.1957, 0.924214, 151.4551], "number2"),  # FX: never round EUR to 1
        ([3.85e-05, 0.1391, 1.1306], None),  # two places would print 0.00: leave as is
        ([3541.5, 1693.0, 40.2], "number1"),  # one small amount keeps the decimal
        ([743.1, 288.4, 0.7], "number1"),  # one-decimal data stays one decimal
        ([2156.4, 1439.4, 881.5], "integer"),  # all large: no decimals needed
    ],
)
def test_decimal_columns_keep_what_their_smallest_value_needs(
    values: list[float], expected: str | None
) -> None:
    # Act / Assert
    assert infer_format("Quote", values)[0] == expected


@pytest.mark.parametrize("prefix", ["c:", ""])
def test_fix_chart_xml_keeps_axes_without_delete_visible(prefix: str) -> None:
    # Arrange: what openpyxl writes for an area chart it loaded (axis settings lost)
    p = prefix
    xml = (
        f'<{p}catAx><{p}axId val="10"/><{p}scaling><{p}orientation val="minMax"/></{p}scaling>'
        f'<{p}axPos val="l"/></{p}catAx><{p}valAx><{p}axId val="100"/><{p}scaling/>'
        f'<{p}delete val="1"/></{p}valAx>'
    )

    # Act
    fixed = fix_chart_xml(xml)

    # Assert
    assert f'</{p}scaling><{p}delete val="0"/><{p}axPos' in fixed
    assert f'<{p}delete val="1"/>' in fixed  # an axis hidden on purpose stays hidden


def test_restyle_drops_the_old_manual_plot_layout(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    from openpyxl.chart.layout import Layout, ManualLayout

    path = _old_style_chart(tmp_path)
    wb = openpyxl.load_workbook(path)
    wb["S"]._charts[0].layout = Layout(manualLayout=ManualLayout(x=0.1, y=0.1, w=0.5, h=0.5))
    wb.save(path)

    # Act
    design.style_sheet(str(path), "S", settings, preset="charts", overwrite=True)

    # Assert
    plot = _charts_xml(path)[0].split("<plotArea>")[1]
    assert "<manualLayout>" not in plot.split("<barChart>")[0]


def test_the_longest_bars_label_inside_their_end(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    rows: list[list[object]] = [["Market", "Profit"], ["US", 1008], ["China", 340], ["EU", 950]]
    path = _block(tmp_path, rows)

    # Act
    charts.add_chart(str(path), "S", "bar", "A1:B4", settings, overwrite=True)

    # Assert: US and EU (within 85% of the longest) inside, China outside
    labels = _charts_xml(path)[0].split("<dLbls>")[1].split("</dLbls>")[0]
    inside = re.findall(r'<dLbl><idx val="(\d+)"/>.*?<dLblPos val="inEnd"/>', labels)
    assert inside == ["0", "2"]
    assert '<dLblPos val="outEnd"/>' in labels  # the rest stay outside
