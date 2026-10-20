from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

import openpyxl
import pytest

from slim_sheet.adapters.engine.cache import TABLE_CACHE, TableCache, file_key
from slim_sheet.config import Settings
from slim_sheet.core.anomalies import find_anomalies
from slim_sheet.core.clean import clean_table
from slim_sheet.core.diff import diff_grids, diff_keyed, same
from slim_sheet.core.errors import ErrorCode, SlimError
from slim_sheet.core.textparse import parse_date, parse_number
from slim_sheet.interfaces.cli import main
from slim_sheet.interfaces.handlers import Handlers
from slim_sheet.services.anomalies import anomalies
from slim_sheet.services.clean import clean
from slim_sheet.services.convert import convert
from slim_sheet.services.diff import diff
from slim_sheet.services.merge import merge
from slim_sheet.services.query import query
from slim_sheet.services.undo import undo
from slim_sheet.services.write import WriteRequest, write


@pytest.fixture
def dirty_xlsx(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Raw"
    ws.append(["Invoice ID", "Customer", "Amount", "Date", "Region"])
    rows = [
        ["INV-1", "  Alice ", "1,200.50", "03/01/2026", "North"],
        ["INV-2", "Bob​", "(300)", "15/01/2026", "north"],
        ["INV-3", "Carol   Ann", "N/A", "2026-01-20", "South"],
        ["INV-3", "Carol   Ann", "N/A", "2026-01-20", "South"],
        [None, None, None, None, None],
        ["INV-5", "Dan", "99", "31/01/2026", "South"],
        ["INV-6", "Eve", "101", "01/02/2026", "South"],
        ["INV-7", "Fay", "100", "02/02/2026", "North"],
        ["INV-8", "Gus", "98", "03/02/2026", "North"],
        ["INV-9", "Hal", "100000", "04/02/2026", "North"],
    ]
    for row in rows:
        ws.append(row)
    path = tmp_path / "dirty.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def customers_csv(tmp_path: Path) -> Path:
    path = tmp_path / "customers.csv"
    path.write_text(
        "CustomerID,Name,Tier\n10000,Ann,gold\n10001,Ben,silver\n10001,Ben2,silver\n99999,Zed,bronze\n",
        encoding="utf-8",
    )
    return path


# --- textparse -----------------------------------------------------------------------


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("1,200.50", 1200.5),
        ("1.200,50", 1200.5),
        ("(300)", -300),
        ("12%", 0.12),
        ("$ 1 200", 1200),
        ("1.234.567", 1234567),
        ("0901234567", None),
        ("abc", None),
        ("", None),
    ],
)
def test_parse_number(text: str, expected: float | None) -> None:
    # Act / Assert
    assert parse_number(text) == expected


@pytest.mark.parametrize(
    ("text", "day_first", "expected"),
    [
        ("2026-01-20", True, dt.date(2026, 1, 20)),
        ("03/01/2026", True, dt.date(2026, 1, 3)),
        ("03/01/2026", False, dt.date(2026, 3, 1)),
        ("3-Jan-26", True, dt.date(2026, 1, 3)),
        ("Jan 3, 2026", True, dt.date(2026, 1, 3)),
        ("2026-01-20 14:30", True, dt.datetime(2026, 1, 20, 14, 30)),
        ("31/02/2026", True, None),
        ("hello", True, None),
    ],
)
def test_parse_date(text: str, day_first: bool, expected: object) -> None:
    # Act / Assert
    assert parse_date(text, day_first) == expected


# --- core clean ----------------------------------------------------------------------


def test_clean_table_runs_ops_in_canonical_order() -> None:
    # Arrange
    columns = ["Name", "Amount", "Empty"]
    body = [[" a ", "1,000", None], ["a", "n/a", None], [None, None, None], ["b", "7", None]]

    # Act
    out = clean_table(columns, body, ["dedupe", "drop_empty_cols", "drop_empty_rows", "trim",
                                      "unify_null", "numbers_from_text"])  # fmt: skip

    # Assert
    assert out.columns == ["Name", "Amount"]
    assert out.rows == [["a", 1000], ["a", None], ["b", 7]]
    assert out.counts["trim"] == 1
    assert out.counts["drop_empty_rows"] == 1


def test_clean_leaves_id_like_text_columns_alone() -> None:
    # Act
    out = clean_table(["Phone"], [["0901"], ["0902"], ["0903"]], ["numbers_from_text"])

    # Assert
    assert out.rows == [["0901"], ["0902"], ["0903"]]
    assert out.counts == {}


def test_clean_rejects_unknown_ops() -> None:
    # Act / Assert
    with pytest.raises(SlimError) as info:
        clean_table(["A"], [["x"]], ["sparkle"])
    assert "trim" in info.value.hint


# --- clean service -------------------------------------------------------------------


def test_clean_preview_writes_nothing(dirty_xlsx: Path, settings: Settings) -> None:
    # Act
    result = clean("dirty.xlsx", ["trim", "unify_null", "numbers_from_text"], settings)

    # Assert
    assert result.output is None
    assert dict(result.counts)["trim"] == 1
    assert any(addr == "Raw!B2" for addr, *_ in result.samples)
    assert not dirty_xlsx.with_name("dirty.clean.xlsx").exists()


def test_clean_apply_writes_typed_copy(dirty_xlsx: Path, settings: Settings) -> None:
    # Arrange
    ops = ["remove_nonprintable", "trim", "collapse_spaces", "unify_null",
           "numbers_from_text", "normalize_dates", "drop_empty_rows", "dedupe"]  # fmt: skip

    # Act
    result = clean("dirty.xlsx", ops, settings, apply=True)

    # Assert
    assert result.output is not None
    assert result.output.target == "dirty.clean.xlsx"
    assert result.rows_after == 8
    rows = list(openpyxl.load_workbook(dirty_xlsx.with_name("dirty.clean.xlsx"))["Raw"].values)
    assert rows[1][:4] == ("INV-1", "Alice", 1200.5, dt.datetime(2026, 1, 3))
    assert rows[2][1:3] == ("Bob", -300)
    assert rows[3][1] == "Carol Ann"


def test_clean_apply_refuses_to_clobber(dirty_xlsx: Path, settings: Settings) -> None:
    # Arrange
    clean("dirty.xlsx", ["trim"], settings, apply=True)

    # Act / Assert
    with pytest.raises(SlimError) as info:
        clean("dirty.xlsx", ["trim"], settings, apply=True)
    assert info.value.code == ErrorCode.WRITE_CONFLICT
    again = clean("dirty.xlsx", ["trim"], settings, apply=True, overwrite=True)
    assert again.output is not None and again.output.backup is not None


# --- anomalies -----------------------------------------------------------------------


def test_find_anomalies_core_kinds() -> None:
    # Arrange
    columns = ["ID", "Amount", "Region"]
    body = [[f"K{i}", 10 + i % 3, "North"] for i in range(12)]
    body[3][0] = "K1"
    body[5][1] = 5000
    body[7][1] = "oops"
    body[8][2] = "north"
    body[9][2] = "North "

    # Act
    kinds = {(a.row, a.col): a.kind for a in find_anomalies(columns, body)}

    # Assert
    assert kinds[(3, 0)] == "duplicate_key"
    assert kinds[(5, 1)] == "outlier"
    assert kinds[(7, 1)] == "type_mismatch"
    assert kinds[(8, 2)] == "variant"
    assert kinds[(9, 2)] == "whitespace"


def test_anomalies_service_addresses(dirty_xlsx: Path, settings: Settings) -> None:
    # Act
    result = anomalies("dirty.xlsx", settings, max_results=50)

    # Assert
    by_address = {f.address: f.kind for f in result.findings}
    assert by_address["Raw!A5"] == "duplicate_key"
    assert by_address["Raw!E3"] == "variant"
    assert by_address["Raw!B2"] == "whitespace"


def test_anomalies_unknown_column(dirty_xlsx: Path, settings: Settings) -> None:
    # Act / Assert
    with pytest.raises(SlimError) as info:
        anomalies("dirty.xlsx", settings, columns=["Nope"])
    assert "Invoice ID" in info.value.hint


# --- merge ---------------------------------------------------------------------------


def test_merge_matches_number_and_text_keys(
    sales_xlsx: Path, customers_csv: Path, settings: Settings
) -> None:
    # Act
    result = merge("sales.xlsx", "customers.csv", ["OrderID=CustomerID"], settings, sheet="Orders")

    # Assert
    assert result.left_rows == 300
    assert result.left_unmatched == 298
    assert result.right_unmatched == 1
    assert result.warnings and "multiplied" in result.warnings[0]
    assert "Name" in result.preview.columns
    assert result.preview.total_rows == 301


def test_merge_anti_and_output(sales_xlsx: Path, customers_csv: Path, settings: Settings) -> None:
    # Act
    result = merge(
        "customers.csv", "sales.xlsx", ["CustomerID=OrderID"], settings,
        how="anti", other_sheet="Orders", out_path="orphans.csv",
    )  # fmt: skip

    # Assert
    assert result.preview.total_rows == 1
    assert result.output is not None
    text = (sales_xlsx.parent / "orphans.csv").read_text(encoding="utf-8")
    assert "99999" in text and "10000" not in text


def test_merge_bad_key_and_how(sales_xlsx: Path, customers_csv: Path, settings: Settings) -> None:
    # Act / Assert
    with pytest.raises(SlimError) as key:
        merge("sales.xlsx", "customers.csv", ["Nope"], settings)
    assert "keys=['A=B']" in key.value.hint
    with pytest.raises(SlimError):
        merge("sales.xlsx", "customers.csv", ["OrderID"], settings, how="cross")


# --- convert -------------------------------------------------------------------------


def test_convert_xlsx_to_csv_defuses_formulas(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["name", "note"])
    ws.append(["a", "placeholder"])
    ws["B2"].value = '=HYPERLINK("http://evil")'
    ws["B2"].data_type = "s"  # text that merely starts with '=', the typical CSV payload
    ws.append(["b", "-5 is fine?"])
    wb.save(tmp_path / "inj.xlsx")

    # Act
    convert("inj.xlsx", "inj.csv", settings)

    # Assert
    text = (tmp_path / "inj.csv").read_text(encoding="utf-8")
    assert "'=HYPERLINK" in text
    assert "'-5 is fine?" in text


def test_convert_round_trips_through_parquet_and_csv(
    sales_xlsx: Path, people_csv: Path, settings: Settings
) -> None:
    # Act
    convert("sales.xlsx", "orders.parquet", settings, sheet="Orders")
    convert("people.csv", "people.xlsx", settings)
    total = query("orders.parquet", 'SELECT count(*), sum("Qty") FROM data', settings)
    people = query("people.xlsx", "SELECT count(*) FROM data", settings)

    # Assert
    assert total.rows[0][0] == 300
    assert people.rows == ((50,),)


def test_convert_workbook_copies_every_sheet(sales_xlsx: Path, settings: Settings) -> None:
    # Act
    result = convert("sales.xlsx", "copy.xlsx", settings)

    # Assert
    wb = openpyxl.load_workbook(sales_xlsx.with_name("copy.xlsx"))
    assert wb.sheetnames == ["Orders", "Notes", "Secret", "Empty"]
    assert wb["Notes"]["D1"].value == "merged"
    assert result.warnings and "values only" in result.warnings[0]


def test_convert_rejects_unknown_suffix(sales_xlsx: Path, settings: Settings) -> None:
    # Act / Assert
    with pytest.raises(SlimError) as info:
        convert("sales.xlsx", "out.docx", settings)
    assert info.value.code == ErrorCode.UNSUPPORTED_FORMAT


# --- diff ----------------------------------------------------------------------------


def test_same_is_loose_about_noise() -> None:
    # Act / Assert
    assert same(1, 1.0)
    assert same(None, "")
    assert same(" a ", "a")
    assert not same(True, 1)
    assert not same("a", "b")


def test_diff_grids_and_keyed() -> None:
    # Act
    cells = list(diff_grids([["a", 1]], [["a", 2], ["new"]]))
    keyed = diff_keyed(
        ["id", "v"], [[1, "x"], [2, "y"]], ["id", "v", "w"], [[2, "z", 0], [3, "q", 0]], "id"
    )

    # Assert
    assert [(d.row, d.col) for d in cells] == [(0, 1), (1, 0)]
    assert keyed.added == ["3"]
    assert keyed.removed == ["1"]
    assert keyed.changed == [("2", "v", "y", "z")]
    assert keyed.columns_added == ["w"]


def test_diff_original_vs_working_copy(sales_xlsx: Path, settings: Settings) -> None:
    # Arrange
    write(WriteRequest("sales.xlsx", "Orders", "E2", [[999]]), settings)
    write(WriteRequest("sales.xlsx", "Extra", "A1", [["hi"]]), settings)

    # Act
    cells = diff("sales.xlsx", "sales.slim.xlsx", settings)
    keyed = diff("sales.xlsx", "sales.slim.xlsx", settings, sheet="Orders", key="OrderID")

    # Assert
    assert cells.total_changes == 1
    assert cells.changes[0] == ("Orders!E2", 12.5, 999)
    assert "sheets added: Extra" in cells.summary
    assert keyed.changes == (("OrderID=10000 / Total", 12.5, 999),)


# --- undo / journal ------------------------------------------------------------------


def test_undo_steps_back_through_writes(sales_xlsx: Path, settings: Settings) -> None:
    # Arrange
    copy = sales_xlsx.with_name("sales.slim.xlsx")
    write(WriteRequest("sales.xlsx", "Orders", "G1", [["one"]]), settings)
    write(WriteRequest("sales.xlsx", "Orders", "G1", [["two"]]), settings)

    # Act / Assert
    first = undo("sales.slim.xlsx", settings)
    assert first.restored_from is not None
    assert openpyxl.load_workbook(copy)["Orders"]["G1"].value == "one"
    second = undo("sales.slim.xlsx", settings)
    assert second.restored_from is None
    assert not copy.exists()
    with pytest.raises(SlimError) as info:
        undo("sales.slim.xlsx", settings)
    assert "Nothing to undo" in info.value.message
    journal = sales_xlsx.parent / ".slim-sheet" / "journal"
    assert len(list(journal.glob("*undone*"))) == 2


def test_undo_restores_overwritten_original(sales_xlsx: Path, settings: Settings) -> None:
    # Arrange
    before = sales_xlsx.read_bytes()
    write(WriteRequest("sales.xlsx", "Orders", "A1", [["changed"]], overwrite=True), settings)

    # Act
    undo("sales.xlsx", settings)

    # Assert
    assert sales_xlsx.read_bytes() == before


# --- cache ---------------------------------------------------------------------------


def test_table_cache_reloads_after_edit(sales_xlsx: Path, settings: Settings) -> None:
    # Arrange
    TABLE_CACHE.clear()
    sql = 'SELECT "Total" FROM "Orders" LIMIT 1'
    assert query("sales.xlsx", sql, settings).rows == ((12.5,),)

    # Act
    wb = openpyxl.load_workbook(sales_xlsx)
    wb["Orders"]["E2"] = 77
    wb.save(sales_xlsx)
    after = query("sales.xlsx", sql, settings)

    # Assert
    assert after.rows == ((77,),) or after.rows == ((77.0,),)


def test_table_cache_lru_eviction(tmp_path: Path) -> None:
    # Arrange
    cache = TableCache(max_files=2)
    paths = []
    for i in range(3):
        p = tmp_path / f"f{i}.csv"
        p.write_text("a\n1\n", encoding="utf-8")
        paths.append(p)

    # Act
    for p in paths:
        cache.put(file_key(p), "data", [])

    # Assert
    assert cache.get(file_key(paths[0]), "data") is None
    assert cache.get(file_key(paths[2]), "data") == []


# --- interfaces ----------------------------------------------------------------------


def test_standard_tier_handlers_render(dirty_xlsx: Path, settings: Settings) -> None:
    # Arrange
    h = Handlers(settings)

    # Act
    preview = h.clean("dirty.xlsx", ["trim"])
    found = h.anomalies("dirty.xlsx")
    converted = h.convert("dirty.xlsx", "dirty.csv")
    reverted = h.undo("dirty.csv")

    # Assert
    assert "PREVIEW" in preview
    assert "duplicate_key" in found
    assert "Revert with undo(path='dirty.csv')" in converted
    merged = h.merge("dirty.xlsx", "dirty.xlsx", ["Invoice ID"], max_tokens=300)
    assert "next: offset" not in merged
    assert "out_path='merged.xlsx'" in merged
    assert "removed it" in reverted


def test_cli_clean_and_diff(
    dirty_xlsx: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    # Arrange
    monkeypatch.setenv("SLIM_SHEET_ROOTS", str(dirty_xlsx.parent))

    # Act
    code = main(["clean", str(dirty_xlsx), "trim", "dedupe", "--apply"])
    out = capsys.readouterr().out
    bad = main(["diff", str(dirty_xlsx), "missing.xlsx"])
    err = json.loads(capsys.readouterr().out)

    # Assert
    assert code == 0
    assert "Wrote dirty.clean.xlsx" in out
    assert bad == 1
    assert err["error"] == "file_not_found"


def test_convert_to_xlsx_strips_control_characters(tmp_path: Path, settings: Settings) -> None:
    # Arrange
    (tmp_path / "ctrl.csv").write_text("name\nbell\x07here\n", encoding="utf-8")

    # Act
    convert("ctrl.csv", "ctrl.xlsx", settings)

    # Assert
    ws = openpyxl.load_workbook(tmp_path / "ctrl.xlsx")["data"]
    assert ws["A2"].value == "bellhere"
