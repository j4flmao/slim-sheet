"""The skill pack's role presets must run as written: each file's expected columns become a sample
workbook, and its build_report arguments are executed against it unchanged."""

from __future__ import annotations

import ast
import re
from pathlib import Path

import openpyxl
import pytest

from slim_sheet.config import Settings
from slim_sheet.core.report_spec import parse_chart, parse_kpi
from slim_sheet.services import report

PACK = Path(__file__).resolve().parents[1] / "skills" / "building-client-reports"
ROLES = sorted((PACK / "reference" / "roles").glob("*.md"))
_PERIOD = re.compile(r"\b(month|week|quarter|year|tháng)\b", re.I)
_TEXT = re.compile(r"\b(region|unit|rep|customer|product|segment|channel|campaign|site|"
                   r"department|location|company|sector|id)\b", re.I)  # fmt: skip


def _columns(text: str) -> list[str]:
    block = text.split("## Expected columns", 1)[1].split("##", 1)[0]
    return re.findall(r"`([^`]+)`", block)


def _arguments(text: str) -> dict[str, object]:
    block = text.split("## build_report arguments", 1)[1].split("```text", 1)[1]
    call = ast.parse(f"f({block.split('```', 1)[0]})", mode="eval").body
    assert isinstance(call, ast.Call)
    return {k.arg: ast.literal_eval(k.value) for k in call.keywords if k.arg}


def _sample(path: Path, columns: list[str]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Data"
    ws.append(columns)
    for row in range(24):
        values: list[object] = []
        for index, name in enumerate(columns):
            if _PERIOD.search(name):
                values.append(f"2026-{row % 12 + 1:02d}")
            elif _TEXT.search(name):
                values.append(f"{name.split()[0]} {row % 5}")
            else:
                values.append(100.0 + 7 * row + 3 * index)
        ws.append(values)
    wb.save(path)


def test_the_pack_has_every_part_the_skill_links_to() -> None:
    # Arrange
    skill = (PACK / "SKILL.md").read_text(encoding="utf-8")

    # Act
    links = re.findall(r"\]\((reference/[^)]+)\)", skill)

    # Assert
    assert skill.startswith("---\nname: building-client-reports\n")
    assert len(skill.splitlines()) < 500  # Anthropic's SKILL.md body limit
    assert all((PACK / link).exists() for link in links), links
    assert len(ROLES) >= 6


@pytest.mark.parametrize("role", ROLES, ids=lambda p: p.stem)
def test_role_preset_runs_as_written(role: Path, tmp_path: Path) -> None:
    # Arrange
    text = role.read_text(encoding="utf-8")
    args = _arguments(text)
    _sample(tmp_path / "sample.xlsx", _columns(text))
    kpis, charts = list(args["kpis"]), list(args["charts"])  # type: ignore[call-overload]

    # Act
    result = report.build_report(
        str(tmp_path / "sample.xlsx"), "Data", str(args["title"]),
        Settings(roots=(tmp_path,)), kpis=kpis, charts=charts,
        subtitle=str(args["subtitle"]), messages=list(args["messages"]),  # type: ignore[call-overload]
        overwrite=True,
    )  # fmt: skip

    # Assert
    assert all(parse_kpi(k) for k in kpis) and all(parse_chart(c) for c in charts)
    assert result.summary.count("\n- KPI ") == len(kpis)
    assert len(openpyxl.load_workbook(tmp_path / "sample.xlsx")["Dashboard"]._charts) == sum(
        not c.startswith("table:") for c in charts
    )


def test_cost_like_kpis_turn_red_when_they_rise(tmp_path: Path) -> None:
    # Arrange
    text = (PACK / "reference" / "roles" / "cfo.md").read_text(encoding="utf-8")
    _sample(tmp_path / "s.xlsx", _columns(text))

    # Act
    report.build_report(str(tmp_path / "s.xlsx"), "Data", "T", Settings(roots=(tmp_path,)),
                        kpis=["Opex=sum:Opex (USD m)|plan:Opex budget (USD m):budget"],
                        overwrite=True)  # fmt: skip

    # Assert
    board = openpyxl.load_workbook(tmp_path / "s.xlsx")["Dashboard"]
    first = next(iter(board.conditional_formatting)).rules[0]
    assert "<=(" in first.formula[0]  # good = at or under budget
