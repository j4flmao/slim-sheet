"""Deterministic fixture workbooks for the agent eval suite (PLAN §12).

Mirrors the shapes used in tests/conftest.py (Orders, messy multi-region Report,
Calc with formulas, a CSV) plus a few extras the tasks need: duplicate invoice
numbers, a sheet name with a space, a broken formula. Nothing random, so the
expected answers in tasks.json stay valid.

Usage: uv run python -m evals.build_fixtures <out_dir>
"""

from __future__ import annotations

import argparse
import datetime as dt
import logging
import sys
from pathlib import Path
from typing import Any

import openpyxl

log = logging.getLogger("evals.fixtures")

ORDER_ROWS = 300
FIXTURE_FILES = ("sales.xlsx", "messy.xlsx", "people.csv")
DUPLICATE_INVOICES = {7: "INV-1003", 19: "INV-1011", 31: "INV-1011"}


def _put(ws: Any, top: int, left: int, rows: list[list[object]]) -> None:
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            ws.cell(row=top + r, column=left + c, value=value)


def _orders(ws: Any) -> None:
    ws.title = "Orders"
    ws.append(["OrderID", "Date", "Region", "Qty", "Total"])
    for i in range(ORDER_ROWS):
        day = dt.datetime(2026, 1, 1) + dt.timedelta(days=i % 90)
        qty = i % 7 + 1
        ws.append([10_000 + i, day, ["North", "South"][i % 2], qty, qty * 12.5])


def _invoices(ws: Any) -> None:
    ws.append(["InvoiceNo", "Customer", "Amount"])
    customers = ["Acme", "Globex", "Initech", "Umbrella"]
    for i in range(40):
        number = DUPLICATE_INVOICES.get(i, f"INV-{1000 + i}")
        ws.append([number, customers[i % 4], 100 + i * 5])


def build_sales(path: Path) -> None:
    wb = openpyxl.Workbook()
    _orders(wb.active)
    _invoices(wb.create_sheet("Invoices"))
    targets = wb.create_sheet("Q1 Targets")
    _put(targets, 1, 1, [["Region", "Target"], ["North", 5000], ["South", 4500]])
    notes = wb.create_sheet("Notes")
    notes["B3"] = "Customer asked for a refund on order 10042"
    notes["B6"] = "Refund approved by finance"
    notes.merge_cells("D1:E1")
    notes["D1"] = "merged"
    hidden = wb.create_sheet("Secret")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "password123"
    wb.save(path)


def build_messy(path: Path) -> None:
    wb = openpyxl.Workbook()
    report = wb.active
    report.title = "Report"
    report["A1"] = "Quarterly report"
    rows: list[list[object]] = [
        ["Item", "Qty", "Price"],
        ["pen", 3, 1.5],
        ["ink", 10, 7.0],
        ["pad", 1, 2.25],
    ]
    _put(report, 3, 1, rows)
    _put(report, 3, 5, [["Code", "Note"], ["A", "refund issued"], ["B", "ok"]])
    calc = wb.create_sheet("Calc")
    calc["A1"], calc["A2"], calc["B1"] = 1, 2, "=SUM(A1:A2)"
    calc["C1"] = "=IFERROR(VLOOKUP(A1,Report!A3:C6,2,FALSE),0)"
    calc["C5"] = "=SUMM(A1:A2)"
    calc["C7"] = "=VLOOKUP(A1,Report!A3:C6,2,FALSE"
    wb.save(path)


def build_people(path: Path) -> None:
    lines = ["name,age,city"] + [f"user{i},{20 + i % 40},City {i % 5}" for i in range(50)]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def build_all(out_dir: Path) -> list[Path]:
    """(Re)create every fixture in out_dir and remove leftovers from earlier writes."""
    out_dir.mkdir(parents=True, exist_ok=True)
    for stale in out_dir.glob("*.slim.*"):
        stale.unlink()
    build_sales(out_dir / "sales.xlsx")
    build_messy(out_dir / "messy.xlsx")
    build_people(out_dir / "people.csv")
    paths = [out_dir / name for name in FIXTURE_FILES]
    log.info("Built %d fixtures in %s", len(paths), out_dir)
    return paths


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("out_dir", type=Path)
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, stream=sys.stderr)
    build_all(args.out_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
