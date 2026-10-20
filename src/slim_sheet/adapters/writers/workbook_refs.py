"""Apply reference rewrites across an openpyxl workbook: formulas, merges, names, data
validation, conditional formatting, tables and autofilters (PLAN §11.5).

The pure rules live in core/shift.py; this module walks openpyxl's object model.
"""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass, field
from typing import Any

from openpyxl.formatting.formatting import ConditionalFormattingList

from slim_sheet.adapters.writers.chart_refs import rename_charts, shift_charts
from slim_sheet.core.errors import ErrorCode, SlimError
from slim_sheet.core.shift import (
    REF_ERROR,
    Shift,
    rename_in_formula,
    shift_address,
    shift_formula,
    shift_sqref,
)


@dataclass
class RefReport:
    formulas: int = 0
    merges: int = 0
    names: int = 0
    rules: int = 0
    charts: int = 0
    broken: list[str] = field(default_factory=list)

    def summary(self) -> str:
        parts = [f"{self.formulas} formula(s)", f"{self.merges} merge(s)", f"{self.names} name(s)"]
        if self.rules:
            parts.append(f"{self.rules} validation/format rule(s)")
        if self.charts:
            parts.append(f"{self.charts} chart reference(s)")
        return "updated " + ", ".join(parts)


def _for_each_formula(
    wb: Any, fn: Callable[[str, str], tuple[str, bool]], report: RefReport
) -> None:
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    new, changed = fn(value, ws.title)
                    if changed:
                        cell.value = new
                        report.formulas += 1
                        if "#REF!" in new and "#REF!" not in value:
                            report.broken.append(f"{ws.title}!{cell.coordinate}")


def _for_each_name(wb: Any, fn: Callable[[str, str], tuple[str, bool]], report: RefReport) -> None:
    scopes = [wb.defined_names, *(getattr(ws, "defined_names", {}) for ws in wb.worksheets)]
    for names in scopes:
        for defined in names.values():
            text = defined.attr_text or ""
            new, changed = fn("=" + text, "")
            if changed:
                defined.attr_text = new[1:]
                report.names += 1


def check_tables(ws: Any, shift: Shift) -> None:
    """Refuse a delete that would remove a whole table (Excel refuses too)."""
    for table in ws.tables.values():
        if shift_address(table.ref, shift) == REF_ERROR:
            raise SlimError(
                ErrorCode.BAD_ARGUMENT,
                f"This would delete all of table '{table.name}' ({table.ref})",
                hint="Delete the table's rows inside its range, or a smaller block.",
            )


def shift_workbook(wb: Any, ws: Any, shift: Shift, apply: Callable[[], None]) -> RefReport:
    """Run `apply` (the openpyxl insert/delete) and repair every reference to `ws`."""
    check_tables(ws, shift)
    report = RefReport()
    merged = [str(r) for r in ws.merged_cells.ranges]
    for rng in merged:
        ws.unmerge_cells(rng)
    apply()
    for rng in merged:
        new = shift_address(rng, shift)
        if new != REF_ERROR:
            ws.merge_cells(new)
            report.merges += int(new != rng)

    def formula_fn(formula: str, home: str) -> tuple[str, bool]:
        return shift_formula(formula, home or ws.title, ws.title, shift)

    _for_each_formula(wb, formula_fn, report)
    _for_each_name(wb, lambda f, _home: shift_formula(f, "", ws.title, shift), report)
    _shift_rules(ws, shift, report)
    report.charts += shift_charts(wb, ws.title, shift)
    for table in ws.tables.values():
        table.ref = shift_address(table.ref, shift)
    if ws.auto_filter.ref:
        new_filter = shift_address(ws.auto_filter.ref, shift)
        ws.auto_filter.ref = None if new_filter == REF_ERROR else new_filter
    return report


def _shift_rules(ws: Any, shift: Shift, report: RefReport) -> None:
    for dv in list(ws.data_validations.dataValidation):
        old = str(dv.sqref)
        new = shift_sqref(old, shift)
        if new != old:
            report.rules += 1
        if new:
            dv.sqref = new
        else:
            ws.data_validations.dataValidation.remove(dv)
    rebuilt = ConditionalFormattingList()
    for cf in ws.conditional_formatting:
        old = str(cf.sqref)
        new = shift_sqref(old, shift)
        report.rules += int(new != old)
        for rule in cf.rules if new else ():
            rebuilt.add(new, rule)
    ws.conditional_formatting = rebuilt


def rename_references(wb: Any, old: str, new: str) -> RefReport:
    report = RefReport()
    _for_each_formula(wb, lambda f, _home: rename_in_formula(f, old, new), report)
    _for_each_name(wb, lambda f, _home: rename_in_formula(f, old, new), report)
    report.charts += rename_charts(wb, old, new)
    return report


def count_references(wb: Any, sheet: str) -> int:
    """Formulas that mention `sheet` (they become #REF! if the sheet is deleted)."""
    probe = RefReport()
    _for_each_formula(wb, lambda f, _home: rename_in_formula(f, sheet, sheet), probe)
    return probe.formulas
