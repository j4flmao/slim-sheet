"""Print-ready page setup for report and table sheets (docs/prompts/design.md §4).

A report someone prints or exports to PDF must not split a table across pages sideways, must
repeat its header row on every page, and must say which page of how many the reader holds.
`print_table` keeps any print area or fit the author already set: it only fills gaps.
"""

from __future__ import annotations

from typing import Any

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

from slim_sheet.core.a1 import format_range
from slim_sheet.core.models import CellRange

LANDSCAPE_FROM_WIDTH = 110  # total column width (characters) that no longer fits portrait A4
MAX_FIT_WIDTH = 300  # wider tables shrunk onto one page width (under ~45%) become unreadable
_MARGINS = PageMargins(left=0.4, right=0.4, top=0.6, bottom=0.6, header=0.3, footer=0.3)
_PAGE_NUMBERS = "Page &P of &N"


def print_dashboard(ws: Any, area: CellRange, title: str) -> None:
    """One landscape page, centred, with the title and page number in the footer."""
    _page(ws, landscape=True, fit_height=1)
    ws.print_area = format_range(area)
    ws.print_options.horizontalCentered = True
    _footer(ws, title)


def print_table(ws: Any, block: CellRange, title: str) -> list[str]:
    """Fit the columns to the page width, repeat the header row, number the pages."""
    if ws.print_area or _fits(ws):
        return []  # the author already decided how this sheet prints
    width = sum(_width(ws, col) for col in range(block.left, block.right + 1))
    landscape = width > LANDSCAPE_FROM_WIDTH
    header = block.top + 1  # no print area: Excel prints the titles and notes around the table
    ws.print_title_rows = f"{header}:{header}"
    _footer(ws, title)
    if width > MAX_FIT_WIDTH:  # keep a readable size: pages run across, first column repeated
        _page(ws, landscape=True, fit_height=None)
        first = get_column_letter(block.left + 1)
        ws.print_title_cols = f"{first}:{first}"
        return ["too wide for one page: prints landscape across pages, header row and first "
                "column repeated"]  # fmt: skip
    _page(ws, landscape=landscape, fit_height=0)
    return [f"prints {'landscape' if landscape else 'portrait'}, fit to width, header repeated"]


def _fits(ws: Any) -> bool:
    props = ws.sheet_properties.pageSetUpPr
    return props is not None and bool(props.fitToPage)


def _width(ws: Any, col: int) -> float:
    width = ws.column_dimensions[get_column_letter(col + 1)].width
    return float(width) if width else 8.43


def _page(ws: Any, landscape: bool, fit_height: int | None) -> None:
    """`fit_height` None prints at 100% (no fitting); 0 fits the width only."""
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_margins = _MARGINS
    if fit_height is None:
        return
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, fit_height


def _footer(ws: Any, title: str) -> None:
    for footer in (ws.oddFooter, ws.evenFooter):
        footer.left.text = title.replace("&", "&&")[:120]
        footer.left.size = 8
        footer.right.text = _PAGE_NUMBERS
        footer.right.size = 8
