"""openpyxl-backed metadata for xlsx/xlsm: formulas, defined names, cell styling.

calamine (the fast reader) only sees cached values; anything about formulas or
formatting comes from here. Loads are lazy and scoped to what the caller needs.
"""

from __future__ import annotations

import warnings
from pathlib import Path
from types import TracebackType
from typing import Any, Self

import openpyxl

from slim_sheet.core.a1 import format_range
from slim_sheet.core.errors import ErrorCode, SlimError
from slim_sheet.core.models import CellRange


def _load(path: Path, read_only: bool) -> Any:
    try:
        with warnings.catch_warnings():  # openpyxl warns about unsupported extensions
            warnings.simplefilter("ignore")
            return openpyxl.load_workbook(path, read_only=read_only, data_only=False)
    except PermissionError as exc:
        raise SlimError(
            ErrorCode.FILE_LOCKED, f"'{path.name}' is locked", hint="Close it in Excel and retry."
        ) from exc
    except Exception as exc:
        raise SlimError(
            ErrorCode.READ_FAILED, f"openpyxl cannot open '{path.name}': {exc}"
        ) from exc


class XlsxMeta:
    """Random-access formula and name lookup over one workbook (full load, on demand)."""

    def __init__(self, path: Path) -> None:
        self._path = path
        self._wb: Any = None

    @property
    def wb(self) -> Any:
        if self._wb is None:
            self._wb = _load(self._path, read_only=False)
        return self._wb

    def sheet_name(self, requested: str) -> str | None:
        return next((n for n in self.wb.sheetnames if n.casefold() == requested.casefold()), None)

    def formula_at(self, sheet: str, row: int, col: int) -> str | None:
        value = self.wb[sheet].cell(row=row + 1, column=col + 1).value
        text = getattr(value, "text", value)  # ArrayFormula keeps its text in .text
        return text if isinstance(text, str) and text.startswith("=") else None

    def defined_names(self) -> list[str]:
        names = list(self.wb.defined_names.keys())
        for ws in self.wb.worksheets:
            names += list(getattr(ws, "defined_names", {}).keys())
        return names

    def name_target(self, name: str) -> tuple[str, str]:
        """(sheet, A1 range) of a single-area defined name."""
        defined = next(
            (d for k, d in self.wb.defined_names.items() if k.casefold() == name.casefold()), None
        )
        if defined is None:
            raise SlimError(
                ErrorCode.BAD_ARGUMENT,
                f"'{name}' is not a defined name",
                hint=f"Names: {', '.join(self.defined_names()) or 'none'}",
            )
        destinations = list(defined.destinations)
        if len(destinations) != 1:
            raise SlimError(ErrorCode.BAD_ARGUMENT, f"Name '{name}' is not a single range")
        sheet, ref = destinations[0]
        return sheet, ref.replace("$", "")

    def close(self) -> None:
        if self._wb is not None:
            self._wb.close()

    def __enter__(self) -> Self:
        return self

    def __exit__(
        self,
        exc_type: type[BaseException] | None,
        exc: BaseException | None,
        tb: TracebackType | None,
    ) -> None:
        self.close()


def sheet_formulas(path: Path) -> dict[str, dict[tuple[int, int], str]]:
    """Every formula in the workbook as {sheet: {(row, col): '=...'}} (streaming read)."""
    wb = _load(path, read_only=True)
    try:
        result: dict[str, dict[tuple[int, int], str]] = {}
        for ws in wb.worksheets:
            found: dict[tuple[int, int], str] = {}
            for row in ws.iter_rows():
                for cell in row:
                    value = getattr(cell, "value", None)
                    text = getattr(value, "text", value)
                    if isinstance(text, str) and text.startswith("="):
                        found[(cell.row - 1, cell.column - 1)] = text
            result[ws.title] = found
        return result
    finally:
        wb.close()


def cell_details(path: Path, sheet: str, row: int, col: int) -> list[tuple[str, str]]:
    """Human-readable facts about one cell: formula, format, font, fill, comment, validation..."""
    with XlsxMeta(path) as meta:
        ws = meta.wb[sheet]
        cell = ws.cell(row=row + 1, column=col + 1)
        facts: list[tuple[str, str]] = []
        formula = meta.formula_at(sheet, row, col)
        if formula:
            facts.append(("formula", formula))
        facts.append(
            (
                "type",
                {
                    "n": "number",
                    "s": "text",
                    "b": "boolean",
                    "d": "date",
                    "f": "formula",
                    "e": "error",
                }.get(cell.data_type, cell.data_type),
            )
        )
        if cell.number_format and cell.number_format != "General":
            facts.append(("number format", cell.number_format))
        facts += _style_facts(cell)
        facts += _context_facts(ws, cell, row, col)
        return facts


def _style_facts(cell: Any) -> list[tuple[str, str]]:
    facts = []
    font = cell.font
    flags = [n for n, on in (("bold", font.b), ("italic", font.i), ("underline", font.u)) if on]
    color = getattr(font.color, "rgb", None) if font.color is not None else None
    if flags or (isinstance(color, str) and color not in ("FF000000", "00000000")):
        facts.append(
            (
                "font",
                ", ".join(flags + ([f"color #{color[-6:]}"] if isinstance(color, str) else [])),
            )
        )
    fill = cell.fill
    if fill is not None and fill.fill_type == "solid":
        facts.append(("fill", f"#{str(fill.fgColor.rgb)[-6:]}"))
    if cell.alignment is not None and (cell.alignment.horizontal or cell.alignment.wrap_text):
        parts = [cell.alignment.horizontal or "", "wrap" if cell.alignment.wrap_text else ""]
        facts.append(("alignment", " ".join(p for p in parts if p)))
    return facts


def _context_facts(ws: Any, cell: Any, row: int, col: int) -> list[tuple[str, str]]:
    facts = []
    if cell.comment is not None:
        facts.append(
            ("comment", f"{cell.comment.text.strip()} (by {cell.comment.author or 'unknown'})")
        )
    if cell.hyperlink is not None:
        facts.append(("hyperlink", str(cell.hyperlink.target or cell.hyperlink.location)))
    for merged in ws.merged_cells.ranges:
        if (
            merged.min_row <= row + 1 <= merged.max_row
            and merged.min_col <= col + 1 <= merged.max_col
        ):
            rng = CellRange(
                merged.min_row - 1, merged.min_col - 1, merged.max_row - 1, merged.max_col - 1
            )
            facts.append(("merged", format_range(rng)))
    for dv in ws.data_validations.dataValidation:
        if cell.coordinate in dv.sqref:
            rule = f"{dv.type}: {dv.formula1}" + (f" .. {dv.formula2}" if dv.formula2 else "")
            facts.append(("validation", rule))
    rules = sum(len(cf.rules) for cf in ws.conditional_formatting if cell.coordinate in cf.sqref)
    if rules:
        facts.append(("conditional formats", str(rules)))
    if ws.protection.sheet and cell.protection.locked:
        facts.append(("protection", "locked (sheet is protected)"))
    return facts
