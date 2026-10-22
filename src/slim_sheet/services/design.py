"""`style_sheet` use-case: make a whole table look professional in one call (PLAN §16.2).

The model names a preset; the server finds the table, styles header/body/total rows from the
brand, infers number formats from headers and values, sizes columns and sets the sheet view,
then reports what it inferred so a model can correct one column in a single follow-up call.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Any

from openpyxl.styles import PatternFill

from slim_sheet.adapters.brand import brand_for, resolve_brand
from slim_sheet.adapters.writers.theme import apply_brand, body_font, style_cell
from slim_sheet.adapters.writers.xlsx_writer import freeze_panes
from slim_sheet.config import Settings
from slim_sheet.core.a1 import format_range as a1_range
from slim_sheet.core.a1 import index_to_col, parse_range
from slim_sheet.core.design import FORMAT_NAMES, Brand, RoleStyle, formats, role_style
from slim_sheet.core.errors import ErrorCode, SlimError
from slim_sheet.core.models import CellRange, OutputResult
from slim_sheet.core.report_spec import infer_format
from slim_sheet.core.styles import choice
from slim_sheet.services.editing import Change, edit_workbook, get_sheet
from slim_sheet.services.styling import MAX_STYLED_CELLS, autofit_columns

PRESETS = ("report", "data", "finance", "striped", "minimal")
_TOTAL_WORDS = ("total", "grand total", "sum", "tổng")


@dataclass(frozen=True, slots=True)
class TableBlock:
    rng: CellRange
    headers: tuple[str, ...]
    total_row: int | None  # 0-based sheet row of a detected/declared total row


def find_table(ws: Any, range_ref: str | None, total: bool | None) -> TableBlock:
    """The block to style: `range_ref`, or the sheet's used range without leading blank rows."""
    rng = parse_range(range_ref) if range_ref else _used(ws)
    if rng.n_rows * rng.n_cols > MAX_STYLED_CELLS:
        raise SlimError(
            ErrorCode.TOO_LARGE, f"{a1_range(rng)} has more than {MAX_STYLED_CELLS} cells"
        )
    headers = tuple(
        str(ws.cell(row=rng.top + 1, column=c + 1).value or "").strip()
        for c in range(rng.left, rng.right + 1)
    )
    last = rng.bottom
    detected = rng.n_rows > 2 and _is_total(ws, last, rng)
    use_total = detected if total is None else total and rng.n_rows > 2
    return TableBlock(rng, headers, last if use_total else None)


def _used(ws: Any) -> CellRange:
    if ws.max_row == 1 and ws.max_column == 1 and ws.cell(1, 1).value is None:
        raise SlimError(
            ErrorCode.BAD_ARGUMENT,
            f"Sheet '{ws.title}' is empty",
            hint=f"write(path, sheet='{ws.title}', range='A1', values=[[...]]) first",
        )
    top = ws.min_row
    while top < ws.max_row and all(c.value is None for c in ws[top]):
        top += 1
    return CellRange(top - 1, ws.min_column - 1, ws.max_row - 1, ws.max_column - 1)


def _is_total(ws: Any, row: int, rng: CellRange) -> bool:
    for col in range(rng.left, rng.right + 1):
        value = ws.cell(row=row + 1, column=col + 1).value
        if isinstance(value, str) and value.strip():
            return value.strip().lower().startswith(_TOTAL_WORDS)
    return False


def column_values(ws: Any, block: TableBlock, col: int) -> list[object]:
    last = block.total_row - 1 if block.total_row is not None else block.rng.bottom
    return [ws.cell(row=r + 1, column=col + 1).value for r in range(block.rng.top + 1, last + 1)]


def style_sheet(
    path: str,
    sheet: str,
    settings: Settings,
    range: str | None = None,
    preset: str = "report",
    number_format: str = "auto",
    total_row: bool | None = None,
    brand: str | None = None,
    out_path: str | None = None,
    overwrite: bool = False,
) -> OutputResult:
    kind = choice(preset, PRESETS, "preset")
    fmt_mode = number_format.strip().lower() or "auto"
    if fmt_mode not in ("auto", "none", *FORMAT_NAMES):
        raise SlimError(
            ErrorCode.BAD_ARGUMENT,
            f"Unknown number_format '{number_format}'",
            hint=f"Use auto, none or one of: {', '.join(FORMAT_NAMES)}",
        )
    chosen, explicit = resolve_brand(brand, settings)

    def mutate(wb: Any) -> Change:
        look = brand_for(wb, chosen, explicit)
        ws = get_sheet(wb, sheet)
        block = find_table(ws, range, total_row)
        notes = style_block(ws, block, kind, look, fmt_mode)
        themed = apply_brand(wb, look, force=explicit)
        where = f"{ws.title}!{a1_range(block.rng)}"
        total = f", total row {block.total_row + 1}" if block.total_row is not None else ""
        head = f"styled {where} as {kind} (brand {look.name}{'' if themed else ', kept theme'})"
        return Change(f"{head}: header row {block.rng.top + 1}{total}; " + "; ".join(notes))

    return edit_workbook(path, settings, f"style_sheet {sheet} {kind}", mutate, out_path, overwrite)


def style_block(ws: Any, block: TableBlock, preset: str, brand: Brand, fmt_mode: str) -> list[str]:
    """Style header, body and total of one table; returns the per-column format notes."""
    numeric, notes = _apply_formats(ws, block, brand, fmt_mode)
    _style_header(ws, block, brand, preset, numeric)
    _style_body(ws, block, brand, preset)
    if block.total_row is not None:
        for col in range(block.rng.left, block.rng.right + 1):
            style_cell(
                ws.cell(row=block.total_row + 1, column=col + 1), role_style("total", brand), brand
            )
    autofit_columns(ws, block.rng)
    _sheet_view(ws, block, brand, preset)
    return notes or ["no number columns"]


def _apply_formats(
    ws: Any, block: TableBlock, brand: Brand, fmt_mode: str
) -> tuple[set[int], list[str]]:
    named = formats(brand)
    numeric: set[int] = set()
    notes: list[str] = []
    rows = range(block.rng.top + 1, block.rng.bottom + 1)
    for offset, col in enumerate(range(block.rng.left, block.rng.right + 1)):
        header = block.headers[offset]
        key, why = infer_format(header, column_values(ws, block, col))
        if key is None or fmt_mode == "none":
            continue
        if fmt_mode != "auto" and key not in ("date", "year"):
            key, why = fmt_mode, "requested"
        numeric.add(col)
        for row in rows:
            ws.cell(row=row + 1, column=col + 1).number_format = named[key]
        notes.append(f"{index_to_col(col)} {header or '?'}: {key} ({why})")
    return numeric, notes


def _style_header(ws: Any, block: TableBlock, brand: Brand, preset: str, numeric: set[int]) -> None:
    style = role_style("header", brand)
    if preset == "minimal":
        style = RoleStyle(bold=True, color=brand.ink, bottom="thin", wrap=True, valign="bottom")
    for col in range(block.rng.left, block.rng.right + 1):
        cell = ws.cell(row=block.rng.top + 1, column=col + 1)
        style_cell(cell, style, brand)
        if col in numeric:  # headers align with their data
            style_cell(cell, RoleStyle(align="right"), brand)


def _style_body(ws: Any, block: TableBlock, brand: Brand, preset: str) -> None:
    band = PatternFill("solid", fgColor=f"FF{brand.band}")
    last = block.total_row - 1 if block.total_row is not None else block.rng.bottom
    for index, row in enumerate(range(block.rng.top + 1, last + 1)):
        for col in range(block.rng.left, block.rng.right + 1):
            cell = ws.cell(row=row + 1, column=col + 1)
            body_font(cell, brand)
            if preset == "striped" and index % 2 == 1:
                cell.fill = band


def _sheet_view(ws: Any, block: TableBlock, brand: Brand, preset: str) -> None:
    view = ws.sheet_view
    view.zoomScale = 100
    if preset in ("report", "finance", "striped"):
        view.showGridLines = False
    if brand.profile != "gov" and preset != "minimal":
        freeze_panes(ws, f"A{block.rng.top + 2}")  # rows only; never freeze columns left of it
    if preset == "data" and not ws.tables:
        last = block.total_row - 1 if block.total_row is not None else block.rng.bottom
        ws.auto_filter.ref = a1_range(
            CellRange(block.rng.top, block.rng.left, last, block.rng.right)
        )
