"""`explain_formula` / `validate_formula` use-cases over a literal formula or a stored one."""

from __future__ import annotations

import openpyxl

from slim_sheet.adapters.formats import sniff
from slim_sheet.adapters.sandbox import resolve_path
from slim_sheet.adapters.xlsx_meta import XlsxMeta
from slim_sheet.config import Settings
from slim_sheet.core.a1 import format_range, parse_range
from slim_sheet.core.errors import ErrorCode, SlimError
from slim_sheet.core.formula import FormulaReport, analyze
from slim_sheet.core.models import CellRange, FileFormat
from slim_sheet.core.validate import ValidationReport, validate
from slim_sheet.services.workbook import open_file


def explain_formula(
    settings: Settings,
    formula: str | None = None,
    path: str | None = None,
    cell: str | None = None,
    sheet: str | None = None,
) -> tuple[FormulaReport, str | None]:
    """Returns the report plus the 'Sheet!A1' address when the formula came from a file."""
    if formula:
        return analyze(formula), None
    if not path or not cell:
        raise SlimError(
            ErrorCode.BAD_ARGUMENT,
            "Pass either formula, or path and cell",
            hint="e.g. explain_formula(formula='=SUM(A1:A9)') or (path='x.xlsx', cell='Sheet1!C7')",
        )
    sheet_name, ref = _split_cell(cell, sheet)
    text, address = _read_formula(path, sheet_name, ref, settings)
    return analyze(text), address


def _split_cell(cell: str, sheet: str | None) -> tuple[str | None, str]:
    sheet_part, sep, ref = cell.rpartition("!")
    if not sep:
        return sheet, cell
    return sheet_part.strip("'") or sheet, ref


def _read_formula(path: str, sheet: str | None, ref: str, settings: Settings) -> tuple[str, str]:
    source = resolve_path(path, settings)
    if sniff(source) != FileFormat.XLSX:
        raise SlimError(
            ErrorCode.UNSUPPORTED_FORMAT,
            f"Formulas can be read from xlsx/xlsm only, not '{source.name}'",
            hint="Pass the formula text instead: explain_formula(formula='=...')",
        )
    rng = parse_range(ref)
    wb = openpyxl.load_workbook(source, read_only=True)
    try:
        names = wb.sheetnames
        name = next((n for n in names if sheet and n.casefold() == sheet.casefold()), None)
        if sheet and name is None:
            raise SlimError(
                ErrorCode.SHEET_NOT_FOUND, f"No sheet '{sheet}'", hint=f"Sheets: {', '.join(names)}"
            )
        name = name or names[0]
        value = wb[name].cell(row=rng.top + 1, column=rng.left + 1).value
    finally:
        wb.close()
    address = f"{name}!{format_range(CellRange(rng.top, rng.left, rng.top, rng.left))}"
    if not isinstance(value, str) or not value.startswith("="):
        raise SlimError(
            ErrorCode.BAD_ARGUMENT,
            f"{address} holds a value, not a formula: {value!r}",
            hint="Use read to inspect values; pass a cell that contains a formula.",
        )
    return value, address


def validate_formula(
    settings: Settings,
    formula: str | None = None,
    path: str | None = None,
    cell: str | None = None,
    sheet: str | None = None,
    target: str | None = None,
) -> tuple[ValidationReport, str | None]:
    """Validate literal or stored formula; with `path`, sheets and defined names are checked too."""
    address = None
    if not formula:
        if not path or not cell:
            raise SlimError(
                ErrorCode.BAD_ARGUMENT,
                "Pass formula, or path and cell",
                hint="e.g. validate_formula(formula='=XLOOKUP(A1,B:B,C:C)', target='2016')",
            )
        sheet_name, ref = _split_cell(cell, sheet)
        formula, address = _read_formula(path, sheet_name, ref, settings)
    sheets: list[str] | None = None
    names: list[str] | None = None
    default_target = "365"
    if path:
        with open_file(path, settings) as handle:
            sheets = [n for n, _ in handle.reader.sheets()]
            if handle.format == FileFormat.XLSX:
                with XlsxMeta(handle.path) as meta:
                    names = meta.defined_names()
            if handle.format == FileFormat.XLS:
                default_target = "xls"
    report = validate(formula, target or default_target, sheets, names)
    return report, address
