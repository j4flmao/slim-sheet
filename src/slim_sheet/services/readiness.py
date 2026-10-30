"""Client-readiness audit (`anomalies scope=client`): what a reviewer, an accessibility checker
or a finance standard (ICAEW, FAST, UK Analysis Function, WCAG) would flag before a workbook
goes to a client. Read only; every finding says how to fix it.
"""

from __future__ import annotations

import posixpath
import re
import zipfile
from collections import Counter
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import openpyxl

from slim_sheet.adapters.writers.chart_alt import chart_targets, rel_targets
from slim_sheet.adapters.writers.sparklines import sheet_parts
from slim_sheet.config import Settings
from slim_sheet.core.errors import ErrorCode, SlimError
from slim_sheet.core.models import AnomalyResult, FileFormat, Finding
from slim_sheet.core.report_spec import infer_format
from slim_sheet.services.chart_restyle import chart_kind
from slim_sheet.services.workbook import open_file

_DEFAULT_NAME = re.compile(r"^(sheet|feuil|tabelle|hoja|foglio|planilha|trang tính)\s*\d*$", re.I)
_RED_NEGATIVE = re.compile(r"\[red\]", re.I)
_UNIT_HINT = re.compile(
    r"[(\[][^)\]]+[)\]]|%|\b(usd|eur|gbp|jpy|vnd|cny|inr|bn|m|k|000s?|units?|pts|bps|x"
    r"|tỷ|triệu|nghìn|đồng|đơn vị)\b",
    re.I,
)
_FRAME = re.compile(r"<(?:\w+:)?graphicFrame\b.*?</(?:\w+:)?graphicFrame>", re.DOTALL)
_DESCR = re.compile(r"<(?:\w+:)?cNvPr\b[^>]*\bdescr=\"[^\"]+\"")
_AMOUNT_KEYS = ("integer", "number1", "number2", "currency", "k", "m", "bn")
TITLE_ROWS = 3  # merges in the first rows are titles, not table cells
MAX_LINES, MAX_SLICES = 4, 5


def client_audit(path: str, settings: Settings, max_results: int = 20) -> AnomalyResult:
    source, display = _xlsx(path, settings)
    wb = openpyxl.load_workbook(source)
    try:
        found = [
            *_workbook_checks(wb),
            *(f for ws in wb.worksheets for f in _sheet_checks(ws)),
            *_alt_text_checks(source),
        ]
    finally:
        wb.close()
    counts = tuple(Counter(f.kind for f in found).items())
    shown = tuple(found[: max(1, min(max_results, 200))])
    return AnomalyResult(display, "workbook", "client", len(wb.sheetnames), shown, counts)


def _xlsx(path: str, settings: Settings) -> tuple[Path, str]:
    with open_file(path, settings) as handle:
        if handle.format != FileFormat.XLSX:
            raise SlimError(
                ErrorCode.UNSUPPORTED_FORMAT,
                f"The client audit reads xlsx/xlsm, not {handle.format.value}",
                hint=f"convert(path='{handle.display}', out_path='copy.xlsx') first",
            )
        return handle.path, handle.display


def _workbook_checks(wb: Any) -> Iterator[Finding]:
    calc = wb.calculation
    if calc is not None and str(calc.calcMode or "").lower() == "manual":
        yield Finding(
            "workbook",
            "manual",
            "manual_calc",
            "calculation is Manual: readers see stale numbers; set it to Automatic "
            "(sheet_layout calc_mode='auto')",
        )
    if calc is not None and calc.iterate:
        yield Finding(
            "workbook",
            "on",
            "iterative_calc",
            "iterative calculation hides circular references; turn it off and remove "
            "the circularity",
        )
    for link in getattr(wb, "_external_links", []):
        target = getattr(getattr(link, "file_link", None), "Target", "") or "another file"
        yield Finding(
            "workbook",
            target,
            "external_link",
            "a link to another file breaks when the client opens it; paste the values "
            "or bring the source in",
        )
    if not (wb.properties.title or "").strip():
        yield Finding(
            "workbook",
            None,
            "no_title",
            "no document title: screen readers and tagged PDFs announce the file name "
            "(build_report sets it)",
        )


def _sheet_checks(ws: Any) -> Iterator[Finding]:
    name = ws.title
    if _DEFAULT_NAME.match(name.strip()):
        yield Finding(
            name,
            name,
            "default_sheet_name",
            "rename the sheet to say what it holds (sheet_ops rename)",
        )
    if ws.sheet_state != "visible":
        yield Finding(
            name,
            ws.sheet_state,
            "hidden_sheet",
            "hidden content surprises reviewers (ICAEW, FAST); unhide it or move it out",
        )
    if ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None and not ws._charts:
        yield Finding(name, None, "empty_sheet", "an empty sheet: delete it")
    yield from _merge_checks(ws)
    yield from _format_checks(ws)
    yield from _unit_checks(ws)
    yield from _chart_checks(ws)


def _merge_checks(ws: Any) -> Iterator[Finding]:
    inside = [str(m) for m in ws.merged_cells.ranges if m.min_row > TITLE_ROWS]
    if inside:
        yield Finding(
            f"{ws.title}!{inside[0]}",
            f"{len(inside)} range(s)",
            "merged_cells",
            "merged cells inside a table break sorting, filters and screen readers; "
            "use Center Across Selection instead",
        )


def _format_checks(ws: Any) -> Iterator[Finding]:
    seen: set[str] = set()
    for row in ws.iter_rows():
        for cell in row:
            fmt = str(cell.number_format or "")
            if fmt in seen or not _RED_NEGATIVE.search(fmt):
                continue
            seen.add(fmt)
            negative = fmt.split(";")[1] if ";" in fmt else ""
            if negative and "-" not in negative and "(" not in negative:
                yield Finding(
                    f"{ws.title}!{cell.coordinate}",
                    fmt,
                    "red_only_negative",
                    "negatives shown by colour alone (Excel Accessibility Checker "
                    "error); keep a minus sign or parentheses",
                )


def _unit_checks(ws: Any) -> Iterator[Finding]:
    """Amount columns need a unit: in the header, or once above the table ('Units: USD m')."""
    header = _header_row(ws)
    if header is None or _unit_above(ws, header):
        return
    missing = []
    for cell in ws[header]:
        if not isinstance(cell.value, str) or _UNIT_HINT.search(cell.value):
            continue
        values = [ws.cell(row=r, column=cell.column).value for r in range(header + 1, header + 21)]
        key = infer_format(cell.value, values)[0]
        if key in _AMOUNT_KEYS and _looks_like_money(cell.value):
            missing.append(f"{cell.column_letter} {cell.value}")
    if missing:
        yield Finding(
            f"{ws.title}!{header}:{header}",
            ", ".join(missing[:4]) + (f" +{len(missing) - 4}" if len(missing) > 4 else ""),
            "missing_unit",
            f"{len(missing)} amount column(s) without a unit; put it in the header "
            "('Revenue (USD m)') or one line above the table ('Units: USD m')",
        )


def _unit_above(ws: Any, header: int) -> bool:
    for row in range(1, header):
        for cell in ws[row]:
            if isinstance(cell.value, str) and _UNIT_HINT.search(cell.value):
                return True
    return False


def _looks_like_money(header: str) -> bool:
    words = set(re.findall(r"[^\W\d_]+", header.lower()))
    return bool(
        words
        & {
            "revenue",
            "sales",
            "income",
            "profit",
            "cost",
            "costs",
            "amount",
            "value",
            "price",
            "spend",
            "budget",
            "ebitda",
            "cap",
            "doanh",
            "lợi",
            "chi",
            "vốn",
        }
    )


def _header_row(ws: Any) -> int | None:
    """The first of the top rows whose filled cells are mostly text, with numbers below."""
    for row in range(1, min(ws.max_row, 6) + 1):
        cells = [c.value for c in ws[row] if c.value is not None]
        below = [c.value for c in ws[row + 1] if c.value is not None] if row < ws.max_row else []
        texts = sum(isinstance(v, str) for v in cells)
        numbers = sum(isinstance(v, int | float) and not isinstance(v, bool) for v in below)
        if len(cells) >= 2 and texts >= len(cells) * 0.8 and numbers:
            return row
    return None


def _chart_checks(ws: Any) -> Iterator[Finding]:
    for index, chart in enumerate(ws._charts, start=1):
        where = f"{ws.title} chart {index}"
        name = type(chart).__name__
        if "3D" in name:
            yield Finding(where, name, "chart_3d", "3-D distorts lengths; use a flat chart")
        if chart_kind(chart) is None and "3D" not in name:
            continue
        if len(getattr(chart, "_charts", [])) > 1:
            yield Finding(
                where,
                name,
                "dual_axis",
                "two charts on one plot (dual axes) suggest false correlations; "
                "use two charts or index both series",
            )
        if chart.title is None:
            yield Finding(
                where, name, "chart_no_title", "give the chart a title that states its message"
            )
        series = len(chart.series)
        if chart_kind(chart) in ("line", "area", "scatter") and series > MAX_LINES:
            yield Finding(
                where,
                f"{series} series",
                "too_many_series",
                f"more than {MAX_LINES} lines tangle; highlight one, grey the rest",
            )


def _alt_text_checks(path: Path) -> Iterator[Finding]:
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        owners = _drawing_owners(archive, names)
        for name in sorted(n for n in names if re.match(r"^xl/drawings/drawing\d+\.xml$", n)):
            rels = name.replace("drawings/", "drawings/_rels/") + ".rels"
            if rels not in names or not chart_targets(name, archive.read(rels).decode("utf-8")):
                continue
            frames = _FRAME.findall(archive.read(name).decode("utf-8"))
            missing = sum(1 for f in frames if "chart" in f and not _DESCR.search(f))
            if missing:
                yield Finding(
                    owners.get(name, name),
                    f"{missing} chart(s)",
                    "chart_no_alt_text",
                    "charts without alt text fail the Accessibility Checker; any "
                    "slim-sheet edit adds it from the chart title",
                )


def _drawing_owners(archive: zipfile.ZipFile, names: set[str]) -> dict[str, str]:
    """Drawing part -> the title of the sheet that shows it."""
    owners: dict[str, str] = {}
    for title, part in sheet_parts(archive).items():
        rels = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
        if rels in names:
            rels_xml = archive.read(rels).decode("utf-8")
            for drawing in rel_targets(part, rels_xml, "/drawing").values():
                owners[drawing] = title
    return owners
