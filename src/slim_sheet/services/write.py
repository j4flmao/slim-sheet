"""`write` use-case: non-destructive, journaled cell writes (PLAN §3.2, §8).

Target selection, journaling and atomic saves come from the shared edit session
(services/editing.py), so writes keep slim-sheet charts and follow working-copy rules.
"""

from __future__ import annotations

from collections.abc import Iterator, Sequence
from dataclasses import dataclass
from typing import Any

from openpyxl.worksheet.formula import ArrayFormula

from slim_sheet.config import Settings
from slim_sheet.core.a1 import format_range, index_to_col, parse_range
from slim_sheet.core.errors import ErrorCode, SlimError
from slim_sheet.core.file_formula import is_dynamic, spill_extent, to_file
from slim_sheet.core.guard import check_formula
from slim_sheet.core.models import CellRange, CellValue, WriteResult
from slim_sheet.core.validate import validate
from slim_sheet.services.editing import Change, edit_workbook

MAX_CELLS = 10_000
MAX_FORMULA_WARNINGS = 5
MAX_SPILL_CHECK = 100_000  # cells looked at for one formula's possible spill area


@dataclass(frozen=True, slots=True)
class WriteRequest:
    path: str
    sheet: str
    range: str
    values: Sequence[Sequence[CellValue]]
    out_path: str | None = None
    overwrite: bool = False


def write(request: WriteRequest, settings: Settings) -> WriteResult:
    if not request.sheet.strip():
        raise SlimError(ErrorCode.BAD_ARGUMENT, "sheet is empty", hint="e.g. sheet='Sheet1'")
    area = _area(request.range, request.values)
    warnings = _check_values(request.values, area)
    created: list[bool] = []

    def mutate(wb: Any) -> Change:
        existing = next(
            (t for t in wb.sheetnames if t.casefold() == request.sheet.casefold()), None
        )
        created.append(existing is None)
        ws = wb[existing] if existing else wb.create_sheet(request.sheet)
        _clear_old_spills(ws, area)
        dynamic: list[tuple[str, CellRange]] = []
        for r, row in enumerate(request.values):
            for c, value in enumerate(row):
                cell = ws.cell(row=area.top + r + 1, column=area.left + c + 1)
                cell.value = _stored(value, cell.coordinate)
                if isinstance(cell.value, ArrayFormula):
                    at = CellRange(area.top + r, area.left + c, area.top + r, area.left + c)
                    dynamic.append((str(value), at))
        blocked = _spill_report(ws, area, request.values, dynamic) if existing or dynamic else []
        return Change(
            f"wrote {ws.title}!{format_range(area)}",
            warnings=tuple(blocked),
            dynamic=tuple((ws.title, format_range(anchor)) for _, anchor in dynamic),
        )

    note = f"write {request.sheet}!{format_range(area)}"
    out = edit_workbook(request.path, settings, note, mutate, request.out_path, request.overwrite)
    return WriteResult(
        source=out.source,
        target=out.target,
        sheet=request.sheet,
        written=area,
        cells=sum(len(row) for row in request.values),
        created_sheet=created[0],
        backup=out.backup,
        warnings=(*warnings, *out.warnings),
    )


def _area(range_ref: str, values: Sequence[Sequence[CellValue]]) -> CellRange:
    if not values or not any(values):
        raise SlimError(
            ErrorCode.BAD_ARGUMENT, "values is empty", hint='e.g. values=[["Total", 120]]'
        )
    rows, cols = len(values), max(len(row) for row in values)
    if rows * cols > MAX_CELLS:
        raise SlimError(
            ErrorCode.TOO_LARGE, f"{rows * cols} cells exceeds the {MAX_CELLS}-cell limit"
        )
    asked = parse_range(range_ref)
    anchor = CellRange(asked.top, asked.left, asked.top, asked.left)
    is_anchor = asked == anchor
    if not is_anchor and (asked.n_rows, asked.n_cols) != (rows, cols):
        raise SlimError(
            ErrorCode.BAD_ARGUMENT,
            f"range {format_range(asked)} is {asked.n_rows}x{asked.n_cols} "
            f"but values are {rows}x{cols}",
            hint=f"Pass only the top-left cell, e.g. range='{format_range(anchor)}'",
        )
    return CellRange(asked.top, asked.left, asked.top + rows - 1, asked.left + cols - 1)


def _check_values(values: Sequence[Sequence[CellValue]], area: CellRange) -> list[str]:
    formulas = 0
    problems: list[str] = []
    for r, row in enumerate(values):
        for c, value in enumerate(row):
            if not (isinstance(value, str) and value.startswith("=")):
                continue
            formulas += 1
            check_formula(value, f"{index_to_col(area.left + c)}{area.top + r + 1}")
            report = validate(value)
            if not report.ok and len(problems) < MAX_FORMULA_WARNINGS:
                where = f"{index_to_col(area.left + c)}{area.top + r + 1}"
                problems.append(
                    f"{where} {value[:60]}: {report.errors[0]} (evaluate check_only=true)"
                )
    if not formulas:
        return []
    return [
        *problems,
        f"{formulas} formula(s) written without cached results; "
        "reads show them empty until recalculated in Excel (evaluate computes them offline)",
    ]


def _stored(value: CellValue, coordinate: str) -> Any:
    """Formulas are stored the way Excel saves them; spilling ones as dynamic arrays."""
    if not (isinstance(value, str) and value.startswith("=")):
        return value
    text = to_file(value)
    return ArrayFormula(ref=coordinate, text=text) if is_dynamic(value) else text


def _formula_text(value: Any) -> str:
    text = value.text if isinstance(value, ArrayFormula) else value
    text = str(text or "")
    return text if text.startswith("=") else "=" + text


def _array_formulas(ws: Any, estimate: bool) -> Iterator[tuple[Any, CellRange]]:
    """(cell, spill range) for every array formula on the sheet that covers several cells.

    The range is the one Excel last spilled; a dynamic formula Excel never calculated (one
    slim-sheet wrote) only names its own cell, so its possible extent is estimated instead.
    """
    for cell in list(ws._cells.values()):
        value = cell.value
        if not isinstance(value, ArrayFormula) or not value.ref:
            continue
        spill = parse_range(str(value.ref))
        if estimate and spill.n_rows * spill.n_cols == 1 and is_dynamic(_formula_text(value)):
            spill = spill_extent(_formula_text(value), spill) or spill
        if spill.n_rows * spill.n_cols > 1:
            yield cell, spill


def _clear_old_spills(ws: Any, area: CellRange) -> None:
    """Overwriting a spilling formula clears the values it had spilled, as Excel does."""
    for cell, spill in list(_array_formulas(ws, estimate=False)):
        anchor = CellRange(cell.row - 1, cell.column - 1, cell.row - 1, cell.column - 1)
        if area.intersect(anchor) is None or not is_dynamic(_formula_text(cell.value)):
            continue
        for r in range(spill.top, spill.bottom + 1):
            for c in range(spill.left, spill.right + 1):
                if (r, c) != (anchor.top, anchor.left):
                    ws.cell(row=r + 1, column=c + 1).value = None


def _spill_report(
    ws: Any,
    area: CellRange,
    values: Sequence[Sequence[CellValue]],
    dynamic: list[tuple[str, CellRange]],
) -> list[str]:
    """Warn when a spilling formula has non-empty cells in its way (Excel shows #SPILL!)."""
    found: list[str] = []
    for cell, spill in _array_formulas(ws, estimate=True):
        origin = CellRange(cell.row - 1, cell.column - 1, cell.row - 1, cell.column - 1)
        overlap = spill.intersect(area)
        if overlap is None or area.intersect(origin) is not None:
            continue
        if any(values[r - area.top][c - area.left] not in (None, "")
               for r in range(overlap.top, overlap.bottom + 1)
               for c in range(overlap.left, overlap.right + 1)
               if c - area.left < len(values[r - area.top])):  # fmt: skip
            found.append(
                f"{format_range(overlap)} is inside the spill range {format_range(spill)} of "
                f"{cell.coordinate} ({_formula_text(cell.value)[:50]}): Excel can show #SPILL!"
            )
    for formula, anchor in dynamic:
        blockers = _blockers(ws, formula, anchor)
        if blockers:
            found.append(
                f"{format_range(anchor)} {formula[:50]} can spill over non-empty "
                f"{', '.join(blockers)}: Excel would show #SPILL! (clear them or move it)"
            )
    return found[:MAX_FORMULA_WARNINGS]


def _blockers(ws: Any, formula: str, anchor: CellRange) -> list[str]:
    extent = spill_extent(formula, anchor)
    if extent is None or extent.n_rows * extent.n_cols > MAX_SPILL_CHECK:
        return []
    taken: list[str] = []
    for r in range(extent.top, extent.bottom + 1):
        for c in range(extent.left, extent.right + 1):
            if (r, c) == (anchor.top, anchor.left) or (r + 1, c + 1) not in ws._cells:
                continue
            if ws._cells[(r + 1, c + 1)].value not in (None, ""):
                taken.append(f"{index_to_col(c)}{r + 1}")
    return taken[:3] + (["..."] if len(taken) > 3 else [])
