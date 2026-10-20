"""Full-tier structure tools: create_workbook, sheet_ops, rows_ops, name_ops (PLAN §5.3)."""

from __future__ import annotations

import re
from typing import Any

import openpyxl
from openpyxl.workbook.defined_name import DefinedName

from slim_sheet.adapters import journal
from slim_sheet.adapters.sandbox import resolve_output
from slim_sheet.adapters.writers.chart_refs import charts_reading
from slim_sheet.adapters.writers.workbook_refs import (
    count_references,
    rename_references,
    shift_workbook,
)
from slim_sheet.adapters.writers.xlsx_writer import save_atomic
from slim_sheet.config import Settings
from slim_sheet.core.a1 import col_to_index, index_to_col, parse_range
from slim_sheet.core.errors import ErrorCode, SlimError
from slim_sheet.core.models import OutputResult
from slim_sheet.core.refs import split_sheet, valid_defined_name
from slim_sheet.core.shift import Axis, Shift, quote_sheet
from slim_sheet.services.editing import Change, edit_workbook, get_sheet

SHEET_OPS = ("add", "rename", "delete", "copy", "move", "hide", "unhide")
ROW_OPS = ("insert_rows", "delete_rows", "insert_cols", "delete_cols")
NAME_OPS = ("add", "delete", "list")
MAX_COUNT = 10_000
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _check_sheet_name(wb: Any, name: str | None) -> str:
    if not name or not name.strip():
        raise SlimError(ErrorCode.BAD_ARGUMENT, "new_name is required", hint="e.g. new_name='Q1'")
    name = name.strip()
    if len(name) > 31 or _BAD_SHEET_CHARS.search(name) or name.startswith("'"):
        raise SlimError(
            ErrorCode.BAD_ARGUMENT,
            f"Invalid sheet name '{name}'",
            hint="At most 31 characters, none of [ ] : * ? / \\ and no leading quote.",
        )
    if any(s.casefold() == name.casefold() for s in wb.sheetnames):
        raise SlimError(ErrorCode.BAD_ARGUMENT, f"Sheet '{name}' already exists")
    return name


# --- create_workbook ----------------------------------------------------------------------


def create_workbook(
    path: str, settings: Settings, sheets: list[str] | None = None, overwrite: bool = False
) -> OutputResult:
    target = resolve_output(path, settings)
    if target.suffix.lower() != ".xlsx":
        raise SlimError(
            ErrorCode.BAD_ARGUMENT,
            "create_workbook writes .xlsx files",
            hint="e.g. path='report.xlsx'",
        )
    if target.exists() and not overwrite:
        raise SlimError(
            ErrorCode.WRITE_CONFLICT, f"'{target.name}' already exists", hint="Pass overwrite=true."
        )
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # start empty so every requested name goes through validation
    names = sheets or ["Sheet1"]
    for name in names:
        wb.create_sheet(_check_sheet_name(wb, name))
    backup = journal.record(target, f"create_workbook {', '.join(names)}")
    save_atomic(wb, target)
    summary = f"new workbook with sheet(s) {', '.join(wb.sheetnames)}"
    return OutputResult(target.name, target.name, summary, backup.name if backup else None)


# --- sheet_ops ----------------------------------------------------------------------------


def sheet_ops(
    path: str,
    op: str,
    sheet: str,
    settings: Settings,
    new_name: str | None = None,
    position: int | None = None,
    out_path: str | None = None,
    overwrite: bool = False,
) -> OutputResult:
    if op not in SHEET_OPS:
        raise SlimError(
            ErrorCode.BAD_ARGUMENT, f"Unknown op '{op}'", hint=f"Use one of: {', '.join(SHEET_OPS)}"
        )

    def mutate(wb: Any) -> Change:
        if op == "add":
            return _add_sheet(wb, sheet, position)
        ws = get_sheet(wb, sheet)
        handlers = {
            "rename": lambda: _rename_sheet(wb, ws, new_name),
            "delete": lambda: _delete_sheet(wb, ws),
            "copy": lambda: _copy_sheet(wb, ws, new_name),
            "move": lambda: _move_sheet(wb, ws, position),
            "hide": lambda: _set_visibility(wb, ws, "hidden"),
            "unhide": lambda: _set_visibility(wb, ws, "visible"),
        }
        return handlers[op]()

    return edit_workbook(path, settings, f"sheet_ops {op} {sheet}", mutate, out_path, overwrite)


def _add_sheet(wb: Any, name: str, position: int | None) -> Change:
    title = _check_sheet_name(wb, name)
    index = None if position is None else max(0, position - 1)
    wb.create_sheet(title, index)
    return Change(f"added sheet '{title}' at position {wb.sheetnames.index(title) + 1}")


def _rename_sheet(wb: Any, ws: Any, new_name: str | None) -> Change:
    old = ws.title
    title = _check_sheet_name(wb, new_name)
    ws.title = title
    report = rename_references(wb, old, title)
    return Change(f"renamed '{old}' to '{title}'; {report.summary()}")


def _delete_sheet(wb: Any, ws: Any) -> Change:
    visible = [s for s in wb.worksheets if s.sheet_state == "visible" and s is not ws]
    if not visible:
        raise SlimError(ErrorCode.BAD_ARGUMENT, "A workbook must keep at least one visible sheet")
    title = ws.title
    wb.remove(ws)
    orphaned = charts_reading(wb, title)
    dangling = count_references(wb, title)
    warnings: tuple[str, ...] = ()
    if dangling:
        warnings += (f"{dangling} formula(s) still reference '{title}' and will show #REF!",)
    if orphaned:
        warnings += (f"{orphaned} chart(s) on other sheets read '{title}' and will show no data",)
    return Change(f"deleted sheet '{title}'", warnings)


def _copy_sheet(wb: Any, ws: Any, new_name: str | None) -> Change:
    title = _check_sheet_name(wb, new_name or f"{ws.title} copy")
    copy = wb.copy_worksheet(ws)
    copy.title = title
    note = "charts, images and data validation are not copied (openpyxl limitation)"
    return Change(f"copied '{ws.title}' to '{title}'", (note,))


def _move_sheet(wb: Any, ws: Any, position: int | None) -> Change:
    if position is None or position < 1:
        raise SlimError(ErrorCode.BAD_ARGUMENT, "position (1 = first) is required for move")
    target = min(position, len(wb.sheetnames)) - 1
    wb.move_sheet(ws, offset=target - wb.sheetnames.index(ws.title))
    return Change(f"moved '{ws.title}' to position {target + 1}")


def _set_visibility(wb: Any, ws: Any, state: str) -> Change:
    if state == "hidden":
        others = [s for s in wb.worksheets if s.sheet_state == "visible" and s is not ws]
        if not others:
            raise SlimError(ErrorCode.BAD_ARGUMENT, "Cannot hide the only visible sheet")
        if wb.active is ws:
            wb.active = wb.worksheets.index(others[0])
    ws.sheet_state = state
    return Change(f"sheet '{ws.title}' is now {state}")


# --- rows_ops -----------------------------------------------------------------------------


def rows_ops(
    path: str,
    sheet: str,
    op: str,
    at: str,
    settings: Settings,
    count: int = 1,
    out_path: str | None = None,
    overwrite: bool = False,
) -> OutputResult:
    if op not in ROW_OPS:
        raise SlimError(
            ErrorCode.BAD_ARGUMENT, f"Unknown op '{op}'", hint=f"Use one of: {', '.join(ROW_OPS)}"
        )
    if not 1 <= count <= MAX_COUNT:
        raise SlimError(ErrorCode.BAD_ARGUMENT, f"count must be 1..{MAX_COUNT}")
    axis: Axis = "rows" if op.endswith("rows") else "cols"
    index = _parse_at(at, axis)
    shift = Shift(axis, index, count if op.startswith("insert") else -count)

    def mutate(wb: Any) -> Change:
        ws = get_sheet(wb, sheet)
        method = getattr(ws, op)
        report = shift_workbook(wb, ws, shift, lambda: method(index, count))
        verb = "inserted" if count > 0 and op.startswith("insert") else "deleted"
        where = at.upper() if axis == "cols" else str(index)
        warnings = tuple(f"{cell} now contains #REF!" for cell in report.broken[:10])
        return Change(
            f"{verb} {count} {axis[:-1]}(s) at {where} in '{ws.title}'; {report.summary()}",
            warnings,
        )

    return edit_workbook(path, settings, f"rows_ops {op} {sheet}!{at}", mutate, out_path, overwrite)


def _parse_at(at: str, axis: str) -> int:
    text = str(at).strip().replace("$", "")
    if axis == "rows":
        if not text.isdigit() or int(text) < 1:
            raise SlimError(
                ErrorCode.BAD_ARGUMENT,
                f"at must be a row number for rows, got '{at}'",
                hint="e.g. at='5'",
            )
        return int(text)
    if text.isdigit() and int(text) >= 1:
        return int(text)
    if text.isalpha() and len(text) <= 3:
        return col_to_index(text) + 1
    raise SlimError(
        ErrorCode.BAD_ARGUMENT,
        f"at must be a column letter for cols, got '{at}'",
        hint="e.g. at='C'",
    )


# --- name_ops -----------------------------------------------------------------------------


def list_names(path: str, settings: Settings) -> list[tuple[str, str]]:
    from slim_sheet.adapters.xlsx_meta import XlsxMeta
    from slim_sheet.services.editing import writable_source

    with XlsxMeta(writable_source(path, settings)) as meta:
        return [(name, str(d.attr_text)) for name, d in meta.wb.defined_names.items()]


def name_ops(
    path: str,
    op: str,
    settings: Settings,
    name: str | None = None,
    ref: str | None = None,
    out_path: str | None = None,
    overwrite: bool = False,
) -> OutputResult:
    if op not in ("add", "delete"):
        raise SlimError(
            ErrorCode.BAD_ARGUMENT, f"Unknown op '{op}'", hint=f"Use one of: {', '.join(NAME_OPS)}"
        )
    if not name or not valid_defined_name(name):
        raise SlimError(
            ErrorCode.BAD_ARGUMENT,
            f"Invalid name '{name}'",
            hint="Start with a letter or _, no spaces, and not a cell address (e.g. TaxRate).",
        )

    def mutate(wb: Any) -> Change:
        existing = next((k for k in wb.defined_names if k.casefold() == name.casefold()), None)
        if op == "delete":
            if existing is None:
                raise SlimError(ErrorCode.BAD_ARGUMENT, f"No defined name '{name}'")
            del wb.defined_names[existing]
            return Change(f"deleted name '{existing}'")
        if existing is not None:
            raise SlimError(
                ErrorCode.BAD_ARGUMENT, f"Name '{existing}' exists", hint="Delete it first."
            )
        target = _absolute_ref(wb, ref)
        wb.defined_names[name] = DefinedName(name, attr_text=target)
        return Change(f"defined {name} = {target}")

    return edit_workbook(path, settings, f"name_ops {op} {name}", mutate, out_path, overwrite)


def _absolute_ref(wb: Any, ref: str | None) -> str:
    sheet, address = split_sheet(ref or "")
    if sheet is None or not address:
        raise SlimError(
            ErrorCode.BAD_ARGUMENT,
            "ref must include a sheet",
            hint="e.g. ref='Rates!B2' or 'Data!A1:D50'",
        )
    ws = get_sheet(wb, sheet)
    rng = parse_range(address)  # validates the address
    start = f"${index_to_col(rng.left)}${rng.top + 1}"
    end = f"${index_to_col(rng.right)}${rng.bottom + 1}"
    body = start if start == end else f"{start}:{end}"
    return f"{quote_sheet(ws.title)}!{body}"
