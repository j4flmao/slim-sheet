"""Create a folder of realistic demo files for trying slim-sheet with any agent.

    uv run python scripts/make_demo.py [folder]      # default: ~/Documents/slim-sheet-demo

Files:
- sales.xlsx      Orders (1,000 rows), a messy Report sheet (title row, two side-by-side tables),
                  Calc (formulas without cached values, a defined name), Targets (for merge/diff)
- dirty.xlsx      text numbers, mixed date formats, stray spaces, duplicates, an outlier, "N/A"
- customers.csv   customer master data with a duplicate id and an orphan
- budget_2016.xlsx formulas using XLOOKUP/IFS (for evaluate check_only=true, target='2016')
"""

from __future__ import annotations

import datetime as dt
import random
import sys
from pathlib import Path

import openpyxl
from openpyxl.workbook.defined_name import DefinedName

REGIONS = ["North", "South", "East", "West"]
PRODUCTS = ["Pen", "Notebook", "Ink", "Stapler", "Folder"]


def sales(path: Path, rng: random.Random) -> None:
    wb = openpyxl.Workbook()
    orders = wb.active
    assert orders is not None
    orders.title = "Orders"
    orders.append(
        ["OrderID", "Date", "Region", "Product", "CustomerID", "Qty", "UnitPrice", "Total"]
    )
    for i in range(1000):
        qty, price = rng.randint(1, 20), rng.choice([1.5, 3.25, 7.0, 12.5, 4.75])
        day = dt.date(2026, 1, 1) + dt.timedelta(days=rng.randint(0, 270))
        orders.append(
            [
                10000 + i,
                day,
                rng.choice(REGIONS),
                rng.choice(PRODUCTS),
                500 + rng.randint(0, 59),
                qty,
                price,
                qty * price,
            ]
        )
    report = wb.create_sheet("Report")
    report["A1"] = "Quarterly report - generated"
    for r, row in enumerate(
        [
            ["Quarter", "Revenue", "Cost"],
            ["Q1", 41000, 30500],
            ["Q2", 45250, 31200],
            ["Q3", 39800, 30100],
        ],
        start=3,
    ):
        for c, value in enumerate(row, start=1):
            report.cell(row=r, column=c, value=value)
    for r, row in enumerate(
        [["Owner", "Note"], ["Lan", "refund issued for INV-77"], ["Minh", "ok"]], start=3
    ):
        for c, value in enumerate(row, start=5):
            report.cell(row=r, column=c, value=value)
    calc = wb.create_sheet("Calc")
    calc["A1"], calc["B1"] = "Total revenue", "=SUM(Orders!H2:H1001)"
    calc["A2"], calc["B2"] = "VAT", "=B1*VatRate"
    calc["A3"], calc["B3"] = "Gross", "=B1+B2"
    calc["A4"], calc["B4"] = "Avg order", "=AVERAGE(Orders!H2:H1001)"
    calc["A5"], calc["B5"] = "Broken lookup", "=VLOOKUP(10005,Orders!A:H,9,FALSE)"
    calc["D1"], calc["E1"] = "VAT rate", 0.1
    wb.defined_names["VatRate"] = DefinedName("VatRate", attr_text="Calc!$E$1")
    targets = wb.create_sheet("Targets")
    targets.append(["Region", "Target"])
    for region, target in zip(REGIONS, [24000, 26000, 22000, 25000], strict=True):
        targets.append([region, target])
    wb.save(path)


def dirty(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Raw"
    ws.append(["Invoice ID", "Customer", "Amount", "Date", "Region"])
    rows = [
        ["INV-1", "  Alice ", "1,200.50", "03/01/2026", "North"],
        ["INV-2", "Bob", "(300)", "15/01/2026", "north"],
        ["INV-3", "Carol   Ann", "N/A", "2026-01-20", "South"],
        ["INV-3", "Carol   Ann", "N/A", "2026-01-20", "South"],
        [None, None, None, None, None],
        ["INV-5", "Dũng", "99", "31/01/2026", "South"],
        ["INV-6", "Eve", "101", "01/02/2026", "South "],
        ["INV-7", "Fay", "100", "Feb 2, 2026", "North"],
        ["INV-8", "Gus", "98", "03.02.2026", "North"],
        ["INV-9", "Hà", "100000", "04/02/2026", "North"],
        ["INV-10", "Ivy", "1.234,56", "05/02/2026", "East"],
    ]
    for row in rows:
        ws.append(row)
    wb.save(path)


def customers(path: Path) -> None:
    lines = ["CustomerID,Name,Tier,City"]
    tiers, cities = ["gold", "silver", "bronze"], ["Hanoi", "HCMC", "Da Nang"]
    for i in range(60):
        lines.append(f"{500 + i},Customer {i},{tiers[i % 3]},{cities[i % 3]}")
    lines += ["505,Customer 5 (duplicate),gold,Hanoi", "999,Orphan Co,bronze,Hue"]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def budget(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Budget"
    ws.append(["Dept", "Plan", "Actual", "Status"])
    for dept, plan, actual in [("Sales", 100, 120), ("Ops", 80, 70), ("IT", 60, 61)]:
        ws.append([dept, plan, actual])
    for r in range(2, 5):
        ws[f"D{r}"] = f'=IFS(C{r}>B{r},"over",C{r}=B{r},"on plan",TRUE,"under")'
    ws["F1"], ws["F2"] = "Lookup IT", '=XLOOKUP("IT",A2:A4,C2:C4)'
    wb.save(path)


def main(argv: list[str]) -> int:
    folder = Path(argv[1]) if len(argv) > 1 else Path.home() / "Documents" / "slim-sheet-demo"
    folder.mkdir(parents=True, exist_ok=True)
    rng = random.Random(42)
    sales(folder / "sales.xlsx", rng)
    dirty(folder / "dirty.xlsx")
    customers(folder / "customers.csv")
    budget(folder / "budget_2016.xlsx")
    sys.stdout.write(f"demo files written to {folder}" + "\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
